import json
import tempfile
from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import unquote
from unittest import mock
from unittest.mock import PropertyMock, patch

import httpx
import fitz
from docx import Document
from django.conf import settings
from django.contrib.auth import authenticate
from django.core import mail
from django.core.management import call_command
from django.test import RequestFactory
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from accounts.models import LoginFailure
from analysis.models import (
    AnalysisItem,
    AnalysisJob,
    AnalysisReport,
    ModelVersion,
    PositionRuleInitialization,
    PositionRuleVersion,
    PromptVersion,
    ReportNote,
)
from analysis.services.analyze import analyze_item
from analysis.services.jobs import create_analysis_job
from analysis.services.jobs import AnalysisJobError
from analysis.services.redaction import redact_resume_text
from analysis.services.rules import create_generated_rule
from analysis.services.schema import ReportValidationError, validate_report_payload
from analysis.tasks import execute_analysis_job, execute_position_rule_initialization
from recruitment.models import (
    Application,
    Candidate,
    ExclusionMarker,
    Notification,
    Position,
    ResumeVersion,
    SyncJob,
)
from recruitment.services.deletion import (
    purge_expired_applications,
    soft_delete_application,
)
from recruitment.services.files import attach_standard_pdf, save_resume_bytes
from recruitment.services.parsing import parse_html, parse_resume
from recruitment.services.sync import (
    ensure_filename,
    extract_items,
    file_info,
    position_jd,
    run_sync_job,
    safe_sync_error,
    upsert_application,
    upsert_candidate,
    upsert_positions,
)
from recruitment.tasks import (
    execute_sync_job,
    pull_application_resume,
    refresh_candidate_resume_preview,
)
from recruitment.integrations.italent import ITalentClient
from reviews.models import (
    PositionReviewer,
    ReviewBatch,
    ReviewItem,
    Reviewer,
)
from reviews.emails import public_review_url
from reviews.services import (
    create_review_batch,
    revoke_batch,
    token_for_batch,
    verify_batch_token,
)
from reviews.tasks import send_review_batch
from talent_pool.models import (
    CandidateNote,
    InterviewResultOption,
    TalentInterview,
    TalentMembership,
    TalentTag,
    TalentTagAssignment,
)
from talent_pool.services import (
    TalentPoolError,
    add_candidate,
    purge_removed_memberships,
    recommend_candidate,
)
from tests.fakes import (
    FakeITalentClient,
    FakeModelGateway,
    FakeRuleGateway,
    RecordingRuleGateway,
    docx_resume_bytes,
)


class WorkflowFixtureMixin:
    def setUp(self):
        super().setUp()
        self.media = tempfile.TemporaryDirectory()
        self.media_override = override_settings(MEDIA_ROOT=self.media.name)
        self.media_override.enable()
        self.admin = User.objects.create_user(
            username="admin",
            password="Admin-password-123",
            role=User.Role.ADMIN,
            must_change_password=False,
            is_staff=True,
        )
        self.hr = User.objects.create_user(
            username="hr",
            password="Hr-password-123",
            role=User.Role.HR,
            must_change_password=False,
        )

    def tearDown(self):
        self.media_override.disable()
        self.media.cleanup()
        super().tearDown()

    def authenticated_client(self, user):
        client = Client()
        client.force_login(user)
        session = client.session
        session["auth_session_version"] = user.session_version
        session["last_activity_at"] = timezone.now().timestamp()
        session["login_started_at"] = timezone.now().timestamp()
        session.save()
        return client

    def create_application(self, *, applicant_id="A-100", position=None):
        position = position or Position.objects.create(
            beisen_position_id="P-100",
            name="后端工程师",
            source_jd="负责后端系统研发。",
        )
        candidate = Candidate.objects.create(
            applicant_id=applicant_id,
            name="测试候选人",
            phone="13800138000",
            email="candidate@example.com",
        )
        resume, _ = save_resume_bytes(
            candidate,
            docx_resume_bytes(),
            "resume.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        resume.extracted_text = "测试候选人 13800138000 candidate@example.com " + (
            "八年后端研发经验，负责招聘系统和异步任务。" * 30
        )
        resume.parse_status = ResumeVersion.ParseStatus.SUCCESS
        resume.parse_quality = 100
        resume.save(
            update_fields=["extracted_text", "parse_status", "parse_quality"]
        )
        application = Application.objects.create(
            application_id=f"APP-{applicant_id}",
            candidate=candidate,
            position=position,
            applied_at=timezone.now(),
            current_resume=resume,
        )
        return application

    def create_rule(self, position):
        rule = PositionRuleVersion.objects.create(
            position=position,
            version=1,
            evaluation_jd="负责后端研发。",
            hard_requirements=[{"name": "后端经验"}],
            dimensions=[{"name": "相关经验", "weight": 100}],
            bonus_items=[],
            rating_thresholds={"priority": 80, "review": 60},
            created_by=self.admin,
        )
        rule.publish(self.admin)
        return rule

    def mark_analyzed(self, application):
        rule = application.position.rule_versions.filter(status=PositionRuleVersion.Status.PUBLISHED).first()
        if not rule:
            rule = self.create_rule(application.position)
        job = AnalysisJob.objects.create(
            position=application.position,
            requested_by=self.hr,
            total_count=1,
        )
        return AnalysisItem.objects.create(
            job=job,
            application=application,
            resume_version=application.current_resume,
            rule_version=rule,
            status=AnalysisItem.Status.SUCCESS,
        )


class AccountTests(WorkflowFixtureMixin, TestCase):
    def test_failed_login_locks_account(self):
        request = RequestFactory().post("/", REMOTE_ADDR="10.0.0.8")
        for _ in range(5):
            self.assertIsNone(
                authenticate(
                    request=request,
                    username=self.hr.username,
                    password="wrong-password",
                )
            )
        self.hr.refresh_from_db()
        self.assertTrue(self.hr.is_locked)
        self.assertEqual(LoginFailure.objects.filter(user=self.hr).count(), 5)
        self.assertEqual(
            LoginFailure.objects.filter(user=self.hr).first().source_ip,
            "10.0.0.8",
        )
        self.assertIsNone(
            authenticate(username=self.hr.username, password="Hr-password-123")
        )

    def test_admin_reset_invalidates_existing_sessions(self):
        client = Client()
        client.force_login(self.hr)
        session = client.session
        session["auth_session_version"] = self.hr.session_version
        session.save()
        admin_client = Client()
        admin_client.force_login(self.admin)
        session = admin_client.session
        session["auth_session_version"] = self.admin.session_version
        session.save()
        response = admin_client.post(
            reverse("accounts:user_reset_password", args=[self.hr.pk])
        )
        self.assertContains(response, "临时密码")
        response = client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith(reverse("accounts:login")))

    def test_inactive_session_expires_after_sixty_minutes(self):
        client = self.authenticated_client(self.hr)
        session = client.session
        session["last_activity_at"] = (
            timezone.now() - timedelta(minutes=61)
        ).timestamp()
        session.save()
        response = client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith(reverse("accounts:login")))

    def test_bootstrap_admin_requires_password_and_creates_admin(self):
        with patch.dict("os.environ", {}, clear=True):
            with self.assertRaisesMessage(Exception, "BOOTSTRAP_ADMIN_PASSWORD"):
                call_command("bootstrap_admin")
        output = StringIO()
        with patch.dict(
            "os.environ",
            {
                "BOOTSTRAP_ADMIN_USERNAME": "bootstrap",
                "BOOTSTRAP_ADMIN_PASSWORD": "Bootstrap-password-123",
            },
            clear=True,
        ):
            call_command("bootstrap_admin", stdout=output)
        user = User.objects.get(username="bootstrap")
        self.assertTrue(user.is_system_admin)
        self.assertTrue(user.must_change_password)

    def test_check_deployment_command(self):
        output = StringIO()
        call_command("check_deployment", "--skip-external", stdout=output)
        text = output.getvalue()
        self.assertIn("部署环境预检", text)
        self.assertIn("检查结果汇总", text)


@override_settings(
    ITALENT_RESUME_MODULES=["ApplicantEducation", "ApplicantWorkExperience"]
)
class SyncTests(WorkflowFixtureMixin, TestCase):
    def test_sync_stops_after_cancellation_request(self):
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.RECONCILIATION,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
        )

        class CancellingClient(FakeITalentClient):
            def iter_applicant_ids(self, start, end, mode=1):
                SyncJob.objects.filter(pk=job.pk).update(
                    status=SyncJob.Status.CANCELLATION_REQUESTED
                )
                yield ["A-1"]

        run_sync_job(job, CancellingClient())

        job.refresh_from_db()
        self.assertEqual(job.status, SyncJob.Status.CANCELLED)
        self.assertFalse(Candidate.objects.exists())

    def test_extract_items_treats_null_data_as_empty(self):
        self.assertEqual(extract_items({"code": 200, "data": None}), [])

    def test_historical_only_candidate_skips_resume_download(self):
        class HistoricalClient(FakeITalentClient):
            def __init__(self):
                super().__init__()
                self.downloads = 0

            def get_positions(self, position_ids):
                payload = super().get_positions(position_ids)
                for item in payload["data"]["items"]:
                    item["status"] = 2
                return payload

            def download_file(self, url):
                self.downloads += 1
                return super().download_file(url)

        client = HistoricalClient()
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.RECONCILIATION,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
        )

        run_sync_job(job, client)
        job.refresh_from_db()

        self.assertEqual(client.downloads, 0)
        self.assertEqual(job.status, SyncJob.Status.SUCCESS)
        self.assertEqual(job.metadata["resume_file_candidates"], 0)
        self.assertEqual(job.metadata["resume_file_skipped_historical"], 1)

    def test_real_italent_dynamic_fields_map_to_display_columns(self):
        profile = {
            "applicantId": "A-REAL",
            "fieldValues": [
                {"name": "Name", "value": "张三", "text": "张三"},
                {"name": "Mobile", "value": "13800138000", "text": "13800138000"},
                {"name": "Email", "value": "real@example.com", "text": "real@example.com"},
                {"name": "LastCompany", "value": "示例公司", "text": "示例公司"},
                {"name": "LastSchool", "value": "示例大学", "text": "示例大学"},
            ],
        }
        modules = {
            "Skill": {
                "applicantId": "A-REAL",
                "moduleInfo": [
                    [
                        {"name": "SkillName", "value": "Python", "text": "Python"},
                        {"name": "SkillLevel", "value": "熟练", "text": "熟练"},
                    ]
                ],
            }
        }

        candidate = upsert_candidate(profile, modules)

        self.assertEqual(candidate.name, "张三")
        self.assertEqual(candidate.phone, "13800138000")
        self.assertEqual(candidate.email, "real@example.com")
        self.assertEqual(candidate.current_company, "示例公司")
        self.assertEqual(candidate.school, "示例大学")
        self.assertIn("Python", candidate.skills_text)

    def test_candidate_school_ignores_beisen_placeholder_code(self):
        candidate = upsert_candidate(
            {
                "applicantId": "A-SCHOOL",
                "fieldValues": [
                    {"name": "Name", "value": "学校测试", "text": "学校测试"},
                    {"name": "LastSchool", "value": "-32767", "text": ""},
                    {
                        "name": "OgLastSchool",
                        "value": "福建工程学院",
                        "text": "福建工程学院",
                    },
                ],
            }
        )

        self.assertEqual(candidate.school, "福建工程学院")

    def test_real_italent_position_fields_map_to_position(self):
        count = upsert_positions(
            {
                "data": [
                    {
                        "jobId": "JOB-REAL",
                        "jobTitle": "真实岗位",
                        "jobType": ["社会招聘"],
                        "duty": "负责系统开发。",
                        "require": "熟悉 Python。",
                        "recruitmentStandard": "三年以上经验。",
                        "status": 1,
                    }
                ]
            }
        )

        position = Position.objects.get(beisen_position_id="JOB-REAL")
        self.assertEqual(count, 1)
        self.assertEqual(position.name, "真实岗位")
        self.assertEqual(position.position_type, "社会招聘")
        self.assertIn("岗位职责", position.source_jd)
        self.assertIn("任职要求", position.source_jd)
        self.assertIn("招聘标准", position.source_jd)
        self.assertEqual(position.status, Position.Status.ACTIVE)

    def test_position_jd_ignores_uuid_recruitment_standard(self):
        value = position_jd(
            {
                "duty": "负责算法研发。",
                "recruitmentStandard": "00000000-0000-0000-0000-000000000000",
            }
        )

        self.assertIn("岗位职责", value)
        self.assertNotIn("招聘标准", value)
        self.assertNotIn("00000000-0000-0000-0000-000000000000", value)

    def test_position_jd_does_not_repeat_section_titles(self):
        value = position_jd(
            {
                "duty": "职位介绍。\n\n岗位职责：\n1. 负责算法研发。",
                "require": "任职要求\n1. 本科及以上学历。",
            }
        )

        self.assertEqual(value.count("岗位职责"), 1)
        self.assertEqual(value.count("任职要求"), 1)
        self.assertIn("1. 负责算法研发。", value)
        self.assertIn("1. 本科及以上学历。", value)

    def test_real_italent_application_field_dictionary_maps_to_columns(self):
        candidate = Candidate.objects.create(applicant_id="A-APPLY")

        application = upsert_application(
            candidate,
            {
                "applyId": "APP-REAL",
                "jobId": "JOB-REAL",
                "fieldValues": {
                    "InitialSubmissionDate": "2026-08-01T10:00:00+08:00",
                    "InitialSubmissionChannel": "招聘官网",
                    "RecruitRequirementId": "REQ-REAL",
                },
            },
        )

        self.assertEqual(application.application_id, "APP-REAL")
        self.assertEqual(application.source_channel, "招聘官网")
        self.assertIsNotNone(application.applied_at)
        self.assertTrue(timezone.is_aware(application.applied_at))
        self.assertEqual(application.position.requisition_id, "REQ-REAL")

    def test_application_channel_id_uses_candidate_profile_display_text(self):
        candidate = Candidate.objects.create(
            applicant_id="A-CHANNEL",
            profile={
                "fieldValues": [
                    {
                        "name": "InitialSubmissionChannel",
                        "value": "channel-id",
                        "text": "BOSS直聘",
                    }
                ]
            },
        )

        application = upsert_application(
            candidate,
            {
                "applyId": "APP-CHANNEL",
                "jobId": "JOB-CHANNEL",
                "fieldValues": {
                    "InitialSubmissionChannel": "channel-id",
                },
            },
        )

        self.assertEqual(application.source_channel, "BOSS直聘")

    def test_protocol_relative_resume_url_and_filename_are_preserved(self):
        info = file_info(
            {
                "data": {
                    "downloadUrl": "//dfiles.italent.cn/download/resume/file.pdf?sig=secret",
                    "dfsPath": "/resume/file.pdf",
                }
            }
        )

        self.assertTrue(info["url"].startswith("//dfiles.italent.cn/"))
        self.assertEqual(info["filename"], "file.pdf")
        self.assertEqual(
            ensure_filename("resume", "application/pdf", b"%PDF"),
            "resume.pdf",
        )

    @patch("recruitment.integrations.italent.httpx.Client")
    def test_download_file_normalizes_url_and_follows_redirects_without_bearer(
        self, client_class
    ):
        response = client_class.return_value.__enter__.return_value.get.return_value
        response.content = b"%PDF"
        response.headers = {"content-type": "application/pdf"}
        client = ITalentClient(client=object())

        content, content_type = client.download_file(
            "//dfiles.italent.cn/download/resume/file.pdf?sig=secret"
        )

        client_class.assert_called_once_with(timeout=60, follow_redirects=True)
        client_class.return_value.__enter__.return_value.get.assert_called_once_with(
            "https://dfiles.italent.cn/download/resume/file.pdf?sig=secret"
        )
        self.assertEqual(content, b"%PDF")
        self.assertEqual(content_type, "application/pdf")

    def test_sync_error_redacts_signed_query_values(self):
        error = safe_sync_error(
            "404 for https://dfiles.italent.cn/file.pdf?sig=secret&token=value"
        )

        self.assertNotIn("secret", error)
        self.assertNotIn("value", error)
        self.assertIn("<redacted>", error)

    def test_fake_italent_sync_is_idempotent_and_preserves_multi_position(self):
        client = FakeITalentClient()
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.FULL,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
        )
        run_sync_job(job, client)
        job.refresh_from_db()
        self.assertEqual(job.status, SyncJob.Status.SUCCESS)
        self.assertEqual(Candidate.objects.count(), 1)
        self.assertEqual(Application.objects.count(), 2)
        self.assertEqual(Position.objects.count(), 2)
        candidate = Candidate.objects.get(applicant_id="A-1")
        self.assertIn("ApplicantEducation", candidate.resume_modules)
        self.assertEqual(candidate.resume_versions.count(), 1)
        self.assertTrue(
            Application.objects.filter(current_resume__isnull=False).count(), 2
        )
        second = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.INCREMENTAL,
            window_start=timezone.now() - timedelta(minutes=15),
            window_end=timezone.now(),
        )
        run_sync_job(second, client)
        self.assertEqual(Candidate.objects.count(), 1)
        self.assertEqual(Application.objects.count(), 2)
        self.assertEqual(candidate.resume_versions.count(), 1)

    def test_scheduled_incremental_sync_omits_empty_runs(self):
        class EmptyClient:
            def iter_applicant_ids(self, start, end, time_type):
                return []

            def get_positions(self, position_ids):
                return {}

        # 1. Scheduled run with no changes deletes itself
        scheduled_job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.INCREMENTAL,
            window_start=timezone.now() - timedelta(minutes=15),
            window_end=timezone.now(),
            requested_by=None,
        )
        run_sync_job(scheduled_job, EmptyClient())
        self.assertFalse(SyncJob.objects.filter(pk=scheduled_job.pk).exists())

        # 2. Manual run with no changes is preserved
        manual_job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.MANUAL,
            window_start=timezone.now() - timedelta(days=1),
            window_end=timezone.now(),
            requested_by=self.hr,
        )
        run_sync_job(manual_job, EmptyClient())
        self.assertTrue(SyncJob.objects.filter(pk=manual_job.pk).exists())

    def test_scheduled_background_sync_does_not_show_progress_bar(self):
        SyncJob.objects.create(
            sync_type=SyncJob.SyncType.INCREMENTAL,
            status=SyncJob.Status.RUNNING,
            window_start=timezone.now() - timedelta(minutes=15),
            window_end=timezone.now(),
            requested_by=None,
        )
        client = self.authenticated_client(self.hr)
        response = client.get(reverse("recruitment:sync_jobs"))
        self.assertNotContains(response, "同步或岗位初始化任务正在执行")

    @override_settings(AUTO_GENERATE_INITIAL_RULES=True)
    @patch("recruitment.services.sync.dispatch_task")
    def test_first_sync_queues_then_generates_beisen_based_rule_v0(
        self, dispatch_task
    ):
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.FULL,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
        )

        run_sync_job(job, FakeITalentClient())

        job.refresh_from_db()
        self.assertEqual(job.status, SyncJob.Status.SUCCESS)
        self.assertEqual(job.metadata["initial_rule_tasks_created"], 2)
        initializations = list(
            PositionRuleInitialization.objects.order_by("position_id")
        )
        self.assertEqual(len(initializations), 2)
        self.assertFalse(PositionRuleVersion.objects.exists())
        self.assertEqual(dispatch_task.call_count, 2)
        for initialization in initializations:
            dispatch_task.assert_any_call(
                execute_position_rule_initialization,
                initialization.pk,
            )

        with patch(
            "analysis.services.rules.ModelGateway",
            return_value=FakeRuleGateway(),
        ) as model_gateway:
            for initialization in initializations:
                execute_position_rule_initialization(initialization.pk)

        for position in Position.objects.all():
            rule = position.rule_versions.get()
            self.assertEqual(rule.version, 0)
            self.assertEqual(rule.status, PositionRuleVersion.Status.PUBLISHED)
            self.assertEqual(rule.evaluation_jd, position.source_jd)
        self.assertEqual(model_gateway.call_count, 2)
        self.assertFalse(
            PositionRuleInitialization.objects.exclude(
                status=PositionRuleInitialization.Status.SUCCESS
            ).exists()
        )

    @override_settings(AUTO_GENERATE_INITIAL_RULES=True)
    @patch("recruitment.services.sync.dispatch_task")
    def test_rule_initialization_failure_does_not_change_sync_result(
        self, dispatch_task
    ):
        class FailingRuleGateway:
            def analyze(self, system_prompt, user_prompt):
                raise RuntimeError("model unavailable")

        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.FULL,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
        )
        run_sync_job(job, FakeITalentClient())
        initialization = PositionRuleInitialization.objects.first()

        with patch(
            "analysis.services.rules.ModelGateway",
            return_value=FailingRuleGateway(),
        ):
            execute_position_rule_initialization(initialization.pk)

        initialization.refresh_from_db()
        job.refresh_from_db()
        self.assertEqual(initialization.status, PositionRuleInitialization.Status.FAILED)
        self.assertEqual(initialization.retry_count, 1)
        self.assertTrue(initialization.error_message)
        self.assertEqual(job.status, SyncJob.Status.SUCCESS)
        self.assertEqual(job.failure_count, 0)
        self.assertTrue(
            self.hr.notifications.filter(title="岗位初始规则生成失败").exists()
        )

    def test_missing_application_id_is_recorded_as_failure(self):
        candidate = Candidate.objects.create(applicant_id="A-X")
        from recruitment.services.sync import upsert_application

        with self.assertRaisesMessage(ValueError, "applicationId"):
            upsert_application(candidate, {"positionId": "P-X"})

    def test_italent_client_handles_pagination_and_token_refresh(self):
        state = {"tokens": 0, "pages": 0, "auth": 0, "unstable": 0}

        def handler(request):
            if request.url.path == "/token":
                state["tokens"] += 1
                return httpx.Response(
                    200,
                    json={"access_token": f"token-{state['tokens']}", "expires_in": 7200},
                )
            if request.url.path.endswith("GetApplicantIdsByDate"):
                state["pages"] += 1
                if state["pages"] == 1:
                    return httpx.Response(
                        200,
                        json={
                            "data": {
                                "applicantIds": ["A-1"],
                                "isLastBatch": False,
                                "nextBatchId": "next",
                            }
                        },
                    )
                return httpx.Response(
                    200,
                    json={
                        "data": {
                            "applicantIds": ["A-2"],
                            "isLastBatch": True,
                        }
                    },
                )
            if request.url.path == "/auth-refresh":
                state["auth"] += 1
                if state["auth"] == 1:
                    return httpx.Response(401, json={"message": "expired"})
                return httpx.Response(200, json={"ok": True})
            if request.url.path == "/unstable":
                state["unstable"] += 1
                if state["unstable"] == 1:
                    return httpx.Response(500, json={"message": "temporary"})
                return httpx.Response(200, json={"ok": True})
            return httpx.Response(404, json={"message": "not found"})

        transport = httpx.MockTransport(handler)
        http = httpx.Client(transport=transport, base_url="https://fake.example")
        with override_settings(
            ITALENT_APP_KEY="key",
            ITALENT_APP_SECRET="secret",
            ITALENT_BASE_URL="https://fake.example",
        ):
            client = ITalentClient(client=http)
            pages = list(
                client.iter_applicant_ids(
                    timezone.now() - timedelta(days=1), timezone.now()
                )
            )
            with patch("recruitment.integrations.italent.time.sleep"):
                self.assertEqual(
                    client.request("GET", "/auth-refresh")["ok"],
                    True,
                )
                self.assertEqual(client.request("GET", "/unstable")["ok"], True)
        self.assertEqual(pages, [["A-1"], ["A-2"]])
        self.assertEqual(state["tokens"], 2)
        self.assertEqual(state["unstable"], 2)

    def test_italent_client_retries_transient_token_connection_failure(self):
        state = {"tokens": 0}

        def handler(request):
            if request.url.path == "/token":
                state["tokens"] += 1
                if state["tokens"] == 1:
                    raise httpx.ConnectError(
                        "[WinError 10013] socket access denied",
                        request=request,
                    )
                return httpx.Response(
                    200,
                    json={"access_token": "token", "expires_in": 7200},
                )
            return httpx.Response(200, json={"data": {"items": []}})

        transport = httpx.MockTransport(handler)
        http = httpx.Client(transport=transport, base_url="https://fake.example")
        with override_settings(
            ITALENT_APP_KEY="key",
            ITALENT_APP_SECRET="secret",
            ITALENT_BASE_URL="https://fake.example",
        ):
            client = ITalentClient(client=http)
            with patch("recruitment.integrations.italent.time.sleep"):
                payload = client.request("GET", "/applicants")

        self.assertEqual(payload, {"data": {"items": []}})
        self.assertEqual(state["tokens"], 2)

    def test_italent_client_accepts_unescaped_control_character_in_json(self):
        response = httpx.Response(
            200,
            content=(
                b'{"code":200,"data":{"items":'
                b'[{"fieldValues":[{"name":"Description","value":"before\x0bafter"}]}]}}'
            ),
        )
        with override_settings(
            ITALENT_APP_KEY="key",
            ITALENT_APP_SECRET="secret",
            ITALENT_BASE_URL="https://fake.example",
        ):
            client = ITalentClient(client=object())
            payload = client._decode(response)

        value = payload["data"]["items"][0]["fieldValues"][0]["value"]
        self.assertEqual(value, "before\x0bafter")

    def test_italent_client_uses_bulk_apply_and_job_request_shapes(self):
        requests = []

        def handler(request):
            if request.url.path == "/token":
                return httpx.Response(
                    200,
                    json={"access_token": "token", "expires_in": 7200},
                )
            requests.append((request.url.path, json.loads(request.content)))
            return httpx.Response(200, json={"data": []})

        transport = httpx.MockTransport(handler)
        http = httpx.Client(transport=transport, base_url="https://fake.example")
        with override_settings(
            ITALENT_APP_KEY="key",
            ITALENT_APP_SECRET="secret",
            ITALENT_BASE_URL="https://fake.example",
            ITALENT_APPLICATIONS_ENDPOINT=(
                "/RecruitV6/api/v1/Apply/GetApplyListByApplicantId"
            ),
            ITALENT_POSITIONS_ENDPOINT="/RecruitV6/api/v1/Job/GetJobListByIds",
            ITALENT_APPLICATION_FIELDS=[
                "InitialSubmissionDate",
                "InitialSubmissionChannel",
            ],
        ):
            client = ITalentClient(client=http)
            client.get_applications(["A-1", "A-2"])
            client.get_positions(["P-1", "P-2"])

        self.assertEqual(
            requests,
            [
                (
                    "/RecruitV6/api/v1/Apply/GetApplyListByApplicantId",
                    {
                        "applicantIds": ["A-1", "A-2"],
                        "fieldNames": [
                            "InitialSubmissionDate",
                            "InitialSubmissionChannel",
                        ],
                    },
                ),
                (
                    "/RecruitV6/api/v1/Job/GetJobListByIds",
                    {"jobIds": ["P-1", "P-2"]},
                ),
            ],
        )

    def test_resume_parser_records_low_quality_and_failure(self):
        candidate = Candidate.objects.create(applicant_id="A-PARSE")
        document = Document()
        document.add_paragraph("简短经历")
        buffer = BytesIO()
        document.save(buffer)
        short_resume, _ = save_resume_bytes(
            candidate,
            buffer.getvalue(),
            "short.docx",
        )
        parse_resume(short_resume)
        self.assertEqual(
            short_resume.parse_status, ResumeVersion.ParseStatus.LOW_QUALITY
        )
        unsupported, _ = save_resume_bytes(
            candidate,
            b"plain text",
            "resume.txt",
        )
        parse_resume(unsupported)
        self.assertEqual(
            unsupported.parse_status,
            ResumeVersion.ParseStatus.UNSUPPORTED,
        )

    def test_resume_parser_falls_back_to_standard_pdf(self):
        candidate = Candidate.objects.create(applicant_id="A-PDF-FALLBACK")
        resume, _ = save_resume_bytes(
            candidate,
            b"\x89PNG\r\n\x1a\n",
            "resume.png",
            "image/png",
        )
        document = fitz.open()
        page = document.new_page()
        page.insert_textbox(
            fitz.Rect(36, 36, 560, 800),
            "\n".join(
                f"Software engineering experience item {index}"
                for index in range(30)
            ),
            fontsize=10,
        )
        attach_standard_pdf(resume, document.tobytes())
        document.close()

        parse_resume(resume)

        self.assertEqual(resume.parse_status, ResumeVersion.ParseStatus.SUCCESS)
        self.assertGreaterEqual(len(resume.extracted_text.strip()), 200)
        self.assertIn("标准 PDF", resume.parse_error)

    def test_html_resume_parser_extracts_visible_text(self):
        candidate = Candidate.objects.create(applicant_id="A-HTML")
        resume, _ = save_resume_bytes(
            candidate,
            (
                "<html><head><style>.hidden{display:none}</style></head>"
                "<body><h1>软件工程师</h1><p>八年系统开发经验。</p>"
                "<script>secret()</script></body></html>"
            ).encode("utf-8"),
            "resume.html",
            "text/html",
        )

        parse_resume(resume)

        self.assertIn("软件工程师", resume.extracted_text)
        self.assertIn("八年系统开发经验", resume.extracted_text)
        self.assertNotIn("secret", resume.extracted_text)
        self.assertIn(
            resume.parse_status,
            {
                ResumeVersion.ParseStatus.SUCCESS,
                ResumeVersion.ParseStatus.LOW_QUALITY,
            },
        )

    def test_resume_parser_handles_remote_storage_without_path(self):
        candidate = Candidate.objects.create(applicant_id="A-S3-STORAGE")
        resume, _ = save_resume_bytes(
            candidate,
            (
                "<html><body><h1>远程云存储候选人</h1><p>"
                + "具备五年分布式与云原生架构经验，负责微服务系统重构与容器化部署运维。" * 10
                + "</p></body></html>"
            ).encode("utf-8"),
            "resume.html",
            "text/html",
        )

        with patch.object(
            type(resume.source_file),
            "path",
            new_callable=mock.PropertyMock,
        ) as mock_path:
            mock_path.side_effect = NotImplementedError("Remote storage does not support path")
            parse_resume(resume)

        self.assertEqual(resume.parse_status, ResumeVersion.ParseStatus.SUCCESS)
        self.assertIn("远程云存储候选人", resume.extracted_text)
        self.assertIn("微服务系统重构", resume.extracted_text)


class AnalysisTests(WorkflowFixtureMixin, TestCase):
    @patch("analysis.tasks.analyze_item")
    def test_analysis_stops_before_next_item_after_cancellation(self, analyze_item):
        first = self.create_application(applicant_id="A-CANCEL-1")
        second = self.create_application(
            applicant_id="A-CANCEL-2",
            position=first.position,
        )
        rule = self.create_rule(first.position)
        job = AnalysisJob.objects.create(
            position=first.position,
            requested_by=self.hr,
            total_count=2,
        )
        first_item = AnalysisItem.objects.create(
            job=job,
            application=first,
            resume_version=first.current_resume,
            rule_version=rule,
        )
        second_item = AnalysisItem.objects.create(
            job=job,
            application=second,
            resume_version=second.current_resume,
            rule_version=rule,
        )

        def finish_first(item):
            item.status = AnalysisItem.Status.SUCCESS
            item.finished_at = timezone.now()
            item.save(update_fields=["status", "finished_at"])
            AnalysisJob.objects.filter(pk=job.pk).update(
                status=AnalysisJob.Status.CANCELLATION_REQUESTED
            )

        analyze_item.side_effect = finish_first
        execute_analysis_job(job.pk)

        job.refresh_from_db()
        first_item.refresh_from_db()
        second_item.refresh_from_db()
        self.assertEqual(job.status, AnalysisJob.Status.CANCELLED)
        self.assertEqual(first_item.status, AnalysisItem.Status.SUCCESS)
        self.assertEqual(second_item.status, AnalysisItem.Status.CANCELLED)
        self.assertEqual(analyze_item.call_count, 1)

    @override_settings(
        MODEL_API_KEY="fake",
        MODEL_NAME="fake-model",
        MODEL_INPUT_COST_PER_MILLION=1.5,
        MODEL_OUTPUT_COST_PER_MILLION=2.5,
    )
    def test_analysis_report_versions_and_reuse(self):
        application = self.create_application()
        self.create_rule(application.position)
        PromptVersion.objects.create(
            version="system-v1", content="system", is_active=True
        )
        ModelVersion.objects.create(
            provider="configured",
            name="fake-model",
            version="fake-model",
            is_active=True,
            input_cost_per_million=1.5,
            output_cost_per_million=2.5,
        )
        job = create_analysis_job(
            application.position, [application.pk], self.hr
        )
        report = analyze_item(job.items.get(), FakeModelGateway())
        self.assertEqual(report.score, 86)
        self.assertEqual(report.rating, AnalysisReport.Rating.PRIORITY)
        self.assertEqual(report.estimated_cost, Decimal("0.0022"))
        application.current_resume.refresh_from_db()
        self.assertTrue(application.current_resume.protected)
        reused_job = create_analysis_job(
            application.position, [application.pk], self.hr
        )
        reused_item = reused_job.items.get()
        self.assertEqual(reused_item.status, AnalysisItem.Status.SUCCESS)
        self.assertEqual(reused_item.reused_report, report)
        forced_job = create_analysis_job(
            application.position,
            [application.pk],
            self.hr,
            force_reason="岗位要求已补充",
        )
        self.assertEqual(forced_job.items.get().status, AnalysisItem.Status.QUEUED)

    @override_settings(
        MODEL_API_KEY="fake",
        MODEL_NAME="transaction-test-model",
        MODEL_INPUT_COST_PER_MILLION=0,
        MODEL_OUTPUT_COST_PER_MILLION=0,
    )
    def test_model_call_does_not_hold_database_transaction(self):
        application = self.create_application(applicant_id="A-TRANSACTION")
        self.create_rule(application.position)
        job = create_analysis_job(application.position, [application.pk], self.hr)
        from django.db import connection

        baseline_atomic_depth = len(connection.atomic_blocks)

        class TransactionCheckingGateway:
            def analyze(self, system_prompt, user_prompt):
                from django.db import connection

                self.atomic_depth = len(connection.atomic_blocks)
                return FakeModelGateway().analyze(system_prompt, user_prompt)

        gateway = TransactionCheckingGateway()
        analyze_item(job.items.get(), gateway)

        self.assertEqual(gateway.atomic_depth, baseline_atomic_depth)

    @override_settings(
        MODEL_API_KEY="fake",
        MODEL_NAME="new-fake-model",
        MODEL_INPUT_COST_PER_MILLION=1.5,
        MODEL_OUTPUT_COST_PER_MILLION=2.5,
    )
    def test_first_analysis_accepts_float_cost_settings(self):
        application = self.create_application()
        self.create_rule(application.position)
        job = create_analysis_job(application.position, [application.pk], self.hr)

        report = analyze_item(job.items.get(), FakeModelGateway())

        self.assertEqual(report.estimated_cost, Decimal("0.0022"))

    def test_redaction_and_schema_validation(self):
        text = redact_resume_text(
            "张三 13800138000 zhangsan@example.com", "张三"
        )
        self.assertNotIn("张三", text)
        self.assertNotIn("13800138000", text)
        self.assertNotIn("zhangsan@example.com", text)
        with self.assertRaises(ReportValidationError):
            validate_report_payload({"score": 101})

    def test_publishing_new_rule_archives_previous(self):
        position = Position.objects.create(name="测试岗位")
        first = self.create_rule(position)
        second = PositionRuleVersion.objects.create(
            position=position,
            version=2,
            evaluation_jd="新版",
            created_by=self.admin,
        )
        second.publish(self.admin)
        first.refresh_from_db()
        self.assertEqual(first.status, PositionRuleVersion.Status.ARCHIVED)
        self.assertEqual(second.status, PositionRuleVersion.Status.PUBLISHED)

    def test_ai_can_generate_editable_rule_draft(self):
        position = Position.objects.create(
            name="平台工程师",
            source_jd="负责平台后端研发、系统设计和稳定性建设。",
        )
        rule = create_generated_rule(position, self.admin, FakeRuleGateway())
        self.assertEqual(rule.status, PositionRuleVersion.Status.DRAFT)
        self.assertEqual(sum(item["weight"] for item in rule.dimensions), 100)

    def test_ai_rule_generation_prefers_system_evaluation_jd(self):
        position = Position.objects.create(
            name="采购开发工程师",
            source_jd="北森原始 JD",
            evaluation_jd="北森与招聘文档合并后的系统评估 JD",
        )
        gateway = RecordingRuleGateway()
        rule = create_generated_rule(position, self.admin, gateway)
        prompt = json.loads(gateway.user_prompt)
        self.assertEqual(
            prompt["source_jd"],
            "北森与招聘文档合并后的系统评估 JD",
        )
        self.assertEqual(
            rule.source_jd_snapshot,
            "北森与招聘文档合并后的系统评估 JD",
        )

    @override_settings(
        MODEL_API_KEY="fake",
        MODEL_NAME="fake-model",
        MODEL_INPUT_COST_PER_MILLION=1,
        MODEL_OUTPUT_COST_PER_MILLION=2,
    )
    def test_report_pages_and_exports_render(self):
        application = self.create_application(applicant_id="A-EXPORT")
        self.create_rule(application.position)
        PromptVersion.objects.create(
            version="system-v1", content="system", is_active=True
        )
        ModelVersion.objects.create(
            provider="configured",
            name="fake-model",
            version="fake-model",
            is_active=True,
        )
        job = create_analysis_job(application.position, [application.pk], self.hr)
        report = analyze_item(job.items.get(), FakeModelGateway())
        report.hard_requirement_results = [
            {
                "name": "相关后端经验",
                "result": "信息不足",
                "evidence": "简历提到后端研发，但未说明完整年限。",
                "note": "建议面试核实。",
            }
        ]
        report.dimension_results = [
            {
                "name": "系统设计",
                "weight": 40,
                "score": 31,
                "evidence": "参与高并发系统建设。",
                "assessment": "具备实践经验。",
            }
        ]
        report.strengths = [
            {"item": "工程落地经验", "evidence": "具备线上项目维护经历。"}
        ]
        report.risks = [
            {"item": "规模信息不足", "evidence": "未说明负责系统的实际规模。"}
        ]
        report.missing_information = [
            {"item": "团队规模", "details": "未说明直接协作人数。"}
        ]
        report.interview_focus = [
            {"focus": "核实系统规模", "reason": "判断经验是否匹配岗位要求。"}
        ]
        report.interview_questions = [
            {
                "question": "请介绍最复杂的系统设计问题。",
                "purpose": "验证系统设计深度。",
            }
        ]
        report.save(
            update_fields=[
                "hard_requirement_results",
                "dimension_results",
                "strengths",
                "risks",
                "missing_information",
                "interview_focus",
                "interview_questions",
            ]
        )
        client = self.authenticated_client(self.hr)
        response = client.get(reverse("analysis:report_detail", args=[report.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "相关后端经验")
        self.assertContains(response, "简历提到后端研发")
        self.assertContains(response, "31 / 40")
        self.assertContains(response, "工程落地经验")
        self.assertContains(response, "验证系统设计深度")
        self.assertNotContains(response, "{&#x27;name&#x27;:")
        self.assertNotContains(response, "{&#x27;item&#x27;:")
        self.assertEqual(
            client.get(reverse("analysis:report_pdf", args=[report.pk])).status_code,
            200,
        )
        self.assertEqual(
            client.get(reverse("analysis:job_excel", args=[job.pk])).status_code,
            200,
        )
        self.assertEqual(
            client.get(reverse("analysis:usage")).status_code,
            302,
        )

    def test_batch_limit_and_partial_failure_status(self):
        first = self.create_application(applicant_id="A-BATCH-1")
        second = self.create_application(
            applicant_id="A-BATCH-2", position=first.position
        )
        self.create_rule(first.position)
        extra_ids = [first.pk, second.pk]
        for index in range(19):
            application = Application.objects.create(
                application_id=f"APP-EXTRA-{index}",
                candidate=first.candidate,
                position=first.position,
                current_resume=first.current_resume,
            )
            extra_ids.append(application.pk)
        with self.assertRaises(AnalysisJobError):
            create_analysis_job(first.position, extra_ids, self.hr)
        job = create_analysis_job(
            first.position,
            [first.pk, second.pk],
            self.hr,
        )

        def fake_analyze(item):
            if item.application_id == first.pk:
                item.status = AnalysisItem.Status.SUCCESS
                item.save(update_fields=["status"])
                return None
            raise RuntimeError("fake model failure")

        from analysis.tasks import execute_analysis_job

        with patch("analysis.tasks.analyze_item", side_effect=fake_analyze):
            execute_analysis_job(job.pk)
        job.refresh_from_db()
        self.assertEqual(job.status, AnalysisJob.Status.PARTIAL)
        self.assertEqual(job.success_count, 1)
        self.assertEqual(job.failure_count, 1)

    def test_analysis_job_page_shows_progress_and_auto_refresh(self):
        application = self.create_application(applicant_id="A-PROGRESS")
        self.create_rule(application.position)
        job = create_analysis_job(application.position, [application.pk], self.hr)
        client = self.authenticated_client(self.hr)

        response = client.get(reverse("analysis:job_detail", args=[job.pk]))

        self.assertContains(response, "AI 正在逐份分析简历")
        self.assertContains(response, "data-auto-refresh")
        self.assertContains(response, "刷新进度")
        self.assertContains(response, "0%")

    @patch("analysis.views.dispatch_task")
    def test_start_analysis_dispatches_background_task(self, dispatch_task):
        application = self.create_application(applicant_id="A-DISPATCH")
        self.create_rule(application.position)
        client = self.authenticated_client(self.hr)

        response = client.post(
            reverse("analysis:start", args=[application.position_id]),
            {"application_ids": [application.pk]},
        )

        job = AnalysisJob.objects.get()
        self.assertRedirects(
            response, reverse("analysis:job_detail", args=[job.pk])
        )
        dispatch_task.assert_called_once_with(execute_analysis_job, job.pk)

    def test_type_error_is_shown_as_system_processing_failure(self):
        application = self.create_application(applicant_id="A-SAFE-ERROR")
        self.create_rule(application.position)
        job = create_analysis_job(application.position, [application.pk], self.hr)

        from analysis.tasks import execute_analysis_job

        with patch(
            "analysis.tasks.analyze_item",
            side_effect=TypeError(
                "unsupported operand type(s) for *: 'decimal.Decimal' and 'float'"
            ),
        ):
            execute_analysis_job(job.pk)

        item = job.items.get()
        self.assertEqual(item.status, AnalysisItem.Status.MODEL_ERROR)
        self.assertEqual(
            item.error_message,
            "系统处理模型结果时发生异常，请联系管理员并重新分析。",
        )
        self.assertNotIn("unsupported operand", item.error_message)


class ReviewTests(WorkflowFixtureMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.application = self.create_application()
        self.analysis_item = self.mark_analyzed(self.application)
        self.prompt_version = PromptVersion.objects.create(
            version="review-public-test", content="system", is_active=True
        )
        self.model_version = ModelVersion.objects.create(
            provider="configured",
            name="internal-model-v9",
            version="internal-model-v9",
            is_active=True,
        )
        self.analysis_report = AnalysisReport.objects.create(
            item=self.analysis_item,
            prompt_version=self.prompt_version,
            model_version=self.model_version,
            score=88,
            rating=AnalysisReport.Rating.PRIORITY,
            hard_requirement_results=[
                {
                    "name": "相关项目经验",
                    "result": "满足",
                    "evidence": "送审时报告：具备招聘系统项目经验。",
                }
            ],
            dimension_results=[
                {
                    "name": "岗位经验",
                    "weight": 40,
                    "score": 35,
                    "assessment": "经历与岗位较匹配。",
                }
            ],
            strengths=[{"item": "工程经验完整", "evidence": "有完整交付经历。"}],
            risks=[{"item": "业务规模待确认", "evidence": "简历未说明用户规模。"}],
            missing_information=[{"item": "团队规模", "details": "建议面试补充。"}],
            interview_focus=[{"focus": "核实项目职责", "reason": "确认实际贡献。"}],
            interview_questions=[
                {
                    "question": "请介绍最有代表性的项目。",
                    "purpose": "验证经历深度。",
                }
            ],
            raw_response={"internal_debug": "不应公开的原始模型数据"},
            input_tokens=123,
            output_tokens=45,
            estimated_cost=Decimal("9.9999"),
        )
        ReportNote.objects.create(
            report=self.analysis_report,
            author=self.hr,
            note_type=ReportNote.NoteType.COMMENT,
            content="HR 内部备注不可见",
        )
        self.reviewer = Reviewer.objects.create(
            name="业务负责人", email="reviewer@example.com"
        )
        PositionReviewer.objects.create(
            position=self.application.position, reviewer=self.reviewer
        )

    def test_token_email_submit_and_link_expiry(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
            72,
        )
        token = token_for_batch(batch)
        self.assertTrue(verify_batch_token(batch, token))
        self.assertFalse(verify_batch_token(batch, "invalid"))
        send_review_batch(batch.pk)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn(token, mail.outbox[0].body)
        self.assertTrue(mail.outbox[0].alternatives)
        html_body = mail.outbox[0].alternatives[0][0]
        self.assertIn(f'href="{public_review_url(batch)}"', html_body)
        self.assertIn("智筛招聘", mail.outbox[0].from_email)
        item = batch.items.get()
        url = reverse("reviews:public", args=[batch.public_id, token])
        response = self.client.get(url)
        self.assertContains(response, self.application.candidate.name)
        detail_response = self.client.get(
            reverse(
                "reviews:public_item",
                args=[batch.public_id, token, item.pk],
            )
        )
        self.assertContains(
            detail_response,
            "送审时报告：具备招聘系统项目经验。",
        )
        response = self.client.post(
            url,
            {
                "action": "submit",
                f"decision_{item.pk}": ReviewItem.Decision.REJECTED,
                f"comment_{item.pk}": "",
            },
        )
        self.assertNotContains(response, "必须填写备注")
        self.assertContains(response, "不能继续使用")
        item.refresh_from_db()
        self.assertEqual(item.decision, ReviewItem.Decision.REJECTED)
        batch.refresh_from_db()
        self.assertEqual(batch.status, ReviewBatch.Status.COMPLETED)

    def test_public_review_shows_snapshot_report_without_internal_data(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        review_item = batch.items.get()
        self.assertEqual(review_item.analysis_report, self.analysis_report)

        newer_job = AnalysisJob.objects.create(
            position=self.application.position,
            requested_by=self.hr,
            total_count=1,
        )
        newer_item = AnalysisItem.objects.create(
            job=newer_job,
            application=self.application,
            resume_version=self.application.current_resume,
            rule_version=self.analysis_item.rule_version,
            status=AnalysisItem.Status.SUCCESS,
        )
        AnalysisReport.objects.create(
            item=newer_item,
            prompt_version=self.prompt_version,
            model_version=self.model_version,
            score=42,
            rating=AnalysisReport.Rating.LOW,
            hard_requirement_results=[
                {
                    "name": "重新分析",
                    "result": "待确认",
                    "evidence": "重新分析后的报告不应替换已送审报告。",
                }
            ],
        )

        url = reverse(
            "reviews:public",
            args=[batch.public_id, token_for_batch(batch)],
        )
        response = self.client.get(url)

        self.assertContains(response, self.application.candidate.name)
        self.assertContains(response, "88")
        self.assertContains(response, "匹配分")
        self.assertContains(response, "未审核")
        self.assertContains(response, "审核")
        self.assertNotContains(response, "送审时报告：具备招聘系统项目经验。")

        detail_response = self.client.get(
            reverse(
                "reviews:public_item",
                args=[
                    batch.public_id,
                    token_for_batch(batch),
                    review_item.pk,
                ],
            )
        )

        self.assertContains(detail_response, "AI 分析内容")
        self.assertContains(detail_response, "审核结果与备注")
        self.assertContains(detail_response, "保存审核结果")
        self.assertContains(detail_response, "在线查看简历")
        self.assertContains(detail_response, "返回审核列表")
        self.assertContains(detail_response, "是否保存审核结果？")
        self.assertContains(detail_response, "AI 分析参考")
        self.analysis_report.interview_focus = [
            {
                "topic": "核实项目职责",
                "reason": "确认候选人的实际贡献。",
            }
        ]
        self.analysis_report.save(update_fields=["interview_focus"])
        topic_response = self.client.get(
            reverse(
                "reviews:public_item",
                args=[batch.public_id, token_for_batch(batch), review_item.pk],
            )
        )
        self.assertContains(topic_response, "核实项目职责")
        self.assertNotContains(topic_response, "未提供")
        self.assertContains(detail_response, "优先评估")
        self.assertContains(detail_response, "相关项目经验")
        self.assertContains(detail_response, "35 / 40")
        self.assertContains(detail_response, "工程经验完整")
        self.assertContains(detail_response, "业务规模待确认")
        self.assertContains(detail_response, "团队规模")
        self.assertContains(detail_response, "核实项目职责")
        self.assertContains(detail_response, "请介绍最有代表性的项目")
        self.assertNotContains(
            detail_response,
            "重新分析后的报告不应替换已送审报告",
        )
        self.assertNotContains(detail_response, "HR 内部备注不可见")
        self.assertNotContains(detail_response, "internal-model-v9")
        self.assertNotContains(detail_response, "9.9999")
        self.assertNotContains(detail_response, "不应公开的原始模型数据")
        self.assertEqual(
            self.client.get(
                reverse("reviews:public", args=[batch.public_id, "invalid"])
            ).status_code,
            404,
        )

    def test_public_reviewer_saves_one_candidate_then_submits_from_list(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        item = batch.items.get()
        token = token_for_batch(batch)
        detail_url = reverse(
            "reviews:public_item",
            args=[batch.public_id, token, item.pk],
        )

        response = self.client.post(
            detail_url,
            {
                "decision": ReviewItem.Decision.REJECTED,
                "comment": "岗位经验需要进一步核实",
            },
        )

        self.assertContains(response, "审核结果已保存")
        item.refresh_from_db()
        self.assertEqual(item.decision, ReviewItem.Decision.REJECTED)
        self.assertEqual(item.comment, "岗位经验需要进一步核实")
        self.assertTrue(item.is_draft)
        self.assertContains(
            self.client.get(
                reverse("reviews:public", args=[batch.public_id, token])
            ),
            "草稿 · 不通过",
        )

        response = self.client.post(
            reverse("reviews:public", args=[batch.public_id, token]),
            {"action": "submit"},
        )

        batch.refresh_from_db()
        item.refresh_from_db()
        self.assertEqual(batch.status, ReviewBatch.Status.COMPLETED)
        self.assertFalse(item.is_draft)
        self.assertContains(response, "不能继续使用")

    def test_public_reviewer_can_clear_saved_draft(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        item = batch.items.get()
        item.decision = ReviewItem.Decision.APPROVED
        item.comment = "先保存的内容"
        item.save(update_fields=["decision", "comment"])

        response = self.client.post(
            reverse(
                "reviews:public",
                args=[batch.public_id, token_for_batch(batch)],
            ),
            {"action": "clear"},
        )

        item.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "已清空")
        self.assertEqual(item.decision, ReviewItem.Decision.PENDING)
        self.assertEqual(item.comment, "")
        self.assertTrue(item.is_draft)
        self.assertIsNone(item.submitted_at)

    def test_logged_in_user_can_view_review_results(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        item = batch.items.get()
        item.decision = ReviewItem.Decision.APPROVED
        item.comment = "建议进入下一轮"
        item.is_draft = False
        item.submitted_at = timezone.now()
        item.save(
            update_fields=["decision", "comment", "is_draft", "submitted_at"]
        )
        client = self.authenticated_client(self.hr)

        response = client.get(reverse("reviews:detail", args=[batch.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.application.candidate.name)
        self.assertContains(response, "建议进入下一轮")
        self.assertContains(
            client.get(reverse("reviews:list")),
            reverse("reviews:detail", args=[batch.pk]),
        )

    def test_review_result_only_shows_talent_import_for_submitted_approval(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        item = batch.items.get()
        item.decision = ReviewItem.Decision.APPROVED
        item.is_draft = False
        item.submitted_at = timezone.now()
        item.save(update_fields=["decision", "is_draft", "submitted_at"])
        client = self.authenticated_client(self.hr)

        response = client.get(reverse("reviews:detail", args=[batch.pk]))

        self.assertContains(response, "导入人才库")
        self.assertContains(
            response,
            reverse(
                "reviews:add_approved_to_talent",
                args=[batch.pk, item.pk],
            ),
        )
        item.decision = ReviewItem.Decision.REJECTED
        item.save(update_fields=["decision"])
        response = client.get(reverse("reviews:detail", args=[batch.pk]))
        self.assertNotContains(response, "导入人才库")

    def test_hr_can_import_approved_review_candidate_to_talent_pool(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        item = batch.items.get()
        item.decision = ReviewItem.Decision.APPROVED
        item.is_draft = False
        item.submitted_at = timezone.now()
        item.save(update_fields=["decision", "is_draft", "submitted_at"])
        client = self.authenticated_client(self.hr)

        response = client.post(
            reverse(
                "reviews:add_approved_to_talent",
                args=[batch.pk, item.pk],
            )
        )

        self.assertRedirects(response, reverse("reviews:detail", args=[batch.pk]))
        membership = TalentMembership.objects.get(
            candidate=self.application.candidate
        )
        self.assertEqual(membership.joined_by, self.hr)
        self.assertEqual(membership.position, self.application.position)
        response = client.get(reverse("reviews:detail", args=[batch.pk]))
        self.assertContains(response, "已在人才库")
        self.assertNotContains(response, "导入人才库")

    def test_hr_cannot_import_unapproved_or_draft_review_result(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        item = batch.items.get()
        item.decision = ReviewItem.Decision.APPROVED
        item.is_draft = True
        item.save(update_fields=["decision", "is_draft"])
        client = self.authenticated_client(self.hr)
        url = reverse(
            "reviews:add_approved_to_talent",
            args=[batch.pk, item.pk],
        )

        response = client.post(url, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "只有已提交且审核通过的候选人才能导入人才库")
        self.assertFalse(
            TalentMembership.objects.filter(
                candidate=self.application.candidate
            ).exists()
        )

    def test_clear_rejected_batch_from_review_and_recycle_bin_retention(self):
        app2 = self.create_application(
            applicant_id="APPLICANT-REJECT-1",
            position=self.application.position,
        )
        app2.candidate.name = "未通过候选人甲"
        app2.candidate.save(update_fields=["name"])
        self.mark_analyzed(app2)

        app3 = self.create_application(
            applicant_id="APPLICANT-REJECT-2",
            position=self.application.position,
        )
        app3.candidate.name = "未通过候选人乙"
        app3.candidate.save(update_fields=["name"])
        self.mark_analyzed(app3)

        batch = create_review_batch(
            self.application.position,
            [self.application.pk, app2.pk, app3.pk],
            self.reviewer,
            self.hr,
        )
        item1 = batch.items.get(application=self.application)
        item1.decision = ReviewItem.Decision.APPROVED
        item1.is_draft = False
        item1.submitted_at = timezone.now()
        item1.save(update_fields=["decision", "is_draft", "submitted_at"])

        item2 = batch.items.get(application=app2)
        item2.decision = ReviewItem.Decision.REJECTED
        item2.is_draft = False
        item2.submitted_at = timezone.now()
        item2.save(update_fields=["decision", "is_draft", "submitted_at"])

        item3 = batch.items.get(application=app3)
        item3.decision = ReviewItem.Decision.REJECTED
        item3.is_draft = False
        item3.submitted_at = timezone.now()
        item3.save(update_fields=["decision", "is_draft", "submitted_at"])

        client = self.authenticated_client(self.hr)
        res = client.get(reverse("reviews:detail", args=[batch.pk]))
        self.assertContains(res, "一键清理")
        self.assertContains(res, "一键导入人才库（1人）")

        # Execute batch clear
        clear_res = client.post(
            reverse("reviews:clear_rejected_batch", args=[batch.pk]),
            follow=True,
        )
        self.assertEqual(clear_res.status_code, 200)
        self.assertContains(clear_res, "已清空 2 名未通过候选人的相关数据")

        app2.refresh_from_db()
        app3.refresh_from_db()
        self.application.refresh_from_db()

        self.assertIsNotNone(app2.deleted_at)
        self.assertIsNotNone(app3.deleted_at)
        self.assertIsNone(self.application.deleted_at)
        self.assertEqual(app2.deleted_by, self.hr)
        self.assertIsNotNone(app2.purge_after)

        # In review detail, rejected candidates show recycled status
        res_after = client.get(reverse("reviews:detail", args=[batch.pk]))
        self.assertNotContains(res_after, "一键清理")
        self.assertContains(res_after, "已放入回收站 (3天后清理)")

        # In recycle bin, Admin can see and restore
        admin_client = self.authenticated_client(self.admin)
        recycle_res = admin_client.get(reverse("recruitment:recycle_bin"))
        self.assertContains(recycle_res, "未通过候选人甲")
        self.assertContains(recycle_res, "未通过候选人乙")

    def test_clear_rejected_single_item_from_review(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        item = batch.items.get()
        item.decision = ReviewItem.Decision.REJECTED
        item.is_draft = False
        item.submitted_at = timezone.now()
        item.save(update_fields=["decision", "is_draft", "submitted_at"])

        client = self.authenticated_client(self.hr)
        res = client.get(reverse("reviews:detail", args=[batch.pk]))
        self.assertContains(res, "清空数据")

        clear_res = client.post(
            reverse("reviews:clear_rejected_item", args=[batch.pk, item.pk]),
            follow=True,
        )
        self.assertEqual(clear_res.status_code, 200)
        self.assertContains(clear_res, "的相关数据，已移入回收站")

        self.application.refresh_from_db()
        self.assertIsNotNone(self.application.deleted_at)

    def test_revoke_and_delete_withdrawal(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        revoke_batch(batch, self.hr)
        batch.refresh_from_db()
        self.assertEqual(batch.status, ReviewBatch.Status.REVOKED)
        second = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        soft_delete_application(self.application, self.hr, "不再处理")
        item = second.items.get()
        self.assertEqual(item.decision, ReviewItem.Decision.WITHDRAWN)

    def test_completed_review_can_reopen_with_new_history(self):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        batch.status = ReviewBatch.Status.COMPLETED
        batch.completed_at = timezone.now()
        batch.save(update_fields=["status", "completed_at"])
        client = self.authenticated_client(self.hr)
        response = client.post(reverse("reviews:reopen", args=[batch.pk]))
        self.assertRedirects(response, reverse("reviews:list"))
        self.assertEqual(ReviewBatch.objects.count(), 2)

    @patch("reviews.views.dispatch_task")
    def test_start_review_can_send_to_multiple_selected_reviewers(self, dispatch_task):
        second_reviewer = Reviewer.objects.create(
            name="第二负责人",
            email="second-reviewer@example.com",
        )
        PositionReviewer.objects.create(
            position=self.application.position,
            reviewer=second_reviewer,
        )
        client = self.authenticated_client(self.hr)

        response = client.post(
            reverse("reviews:start", args=[self.application.position.pk]),
            {
                "application_ids": [self.application.pk],
                "reviewer_ids": [self.reviewer.pk, second_reviewer.pk],
                "expiry_hours": "72",
            },
        )

        self.assertRedirects(response, reverse("reviews:list"))
        batches = ReviewBatch.objects.order_by("reviewer_id")
        self.assertEqual(batches.count(), 2)
        self.assertEqual(
            set(batches.values_list("reviewer_id", flat=True)),
            {self.reviewer.pk, second_reviewer.pk},
        )
        self.assertTrue(all(batch.items.count() == 1 for batch in batches))
        self.assertEqual(
            [call.args[0].__name__ for call in dispatch_task.call_args_list],
            ["send_review_batch", "send_review_batch"],
        )

    @patch("reviews.views.dispatch_task")
    def test_resend_marks_existing_batch_as_email_pending(self, dispatch_task):
        batch = create_review_batch(
            self.application.position,
            [self.application.pk],
            self.reviewer,
            self.hr,
        )
        batch.status = ReviewBatch.Status.EMAIL_FAILED
        batch.email_status = ReviewBatch.EmailStatus.FAILED
        batch.save(update_fields=["status", "email_status"])
        client = self.authenticated_client(self.hr)

        response = client.post(reverse("reviews:resend", args=[batch.pk]))

        self.assertRedirects(response, reverse("reviews:list"))
        batch.refresh_from_db()
        self.assertEqual(batch.status, ReviewBatch.Status.EMAIL_PENDING)
        self.assertEqual(batch.email_status, ReviewBatch.EmailStatus.PENDING)
        dispatch_task.assert_called_once_with(send_review_batch, batch.pk)
        list_response = client.get(reverse("reviews:list"))
        self.assertNotContains(list_response, "data-auto-refresh")
        self.assertContains(list_response, "刷新列表")


class TalentPoolTests(WorkflowFixtureMixin, TestCase):
    def test_talent_list_uses_view_operation(self):
        application = self.create_application()
        membership = add_candidate(application.candidate, self.hr)
        client = self.authenticated_client(self.hr)

        response = client.get(reverse("talent_pool:list"))

        self.assertContains(response, "查看")
        self.assertContains(
            response,
            reverse("talent_pool:detail", args=[membership.pk]),
        )
        self.assertNotContains(response, "推荐岗位")

    def test_talent_detail_hides_online_resume_text_and_shows_unified_notes(self):
        application = self.create_application()
        membership = add_candidate(application.candidate, self.hr)
        CandidateNote.objects.create(
            candidate=application.candidate,
            author=self.hr,
            scope=CandidateNote.Scope.GENERAL,
            content="候选人有招聘系统项目经验。",
        )
        client = self.authenticated_client(self.hr)

        response = client.get(
            reverse("talent_pool:detail", args=[membership.pk])
        )

        self.assertNotContains(response, "在线简历内容")
        self.assertNotContains(response, "八年后端研发经验")
        self.assertContains(response, "备注")
        self.assertContains(response, "保存备注")
        self.assertContains(response, "候选人有招聘系统项目经验。")
        self.assertNotContains(response, "Scope:")
        self.assertNotContains(response, "Content:")

        # Test editing the unified note directly
        edit_res = client.post(
            reverse("talent_pool:add_note", args=[membership.pk]),
            {"content": "更新后的统一备注内容"},
        )
        self.assertRedirects(edit_res, reverse("talent_pool:detail", args=[membership.pk]))
        self.assertEqual(CandidateNote.objects.filter(candidate=application.candidate).count(), 1)
        self.assertEqual(
            CandidateNote.objects.filter(candidate=application.candidate).first().content,
            "更新后的统一备注内容",
        )

    def test_talent_detail_uses_latest_valid_resume_with_preview_and_download(self):
        application = self.create_application()
        membership = add_candidate(application.candidate, self.hr)
        latest_resume, _ = save_resume_bytes(
            application.candidate,
            b"\x89PNG\r\n\x1a\n",
            "latest-resume.png",
            "image/png",
        )
        document = fitz.open()
        page = document.new_page()
        page.insert_textbox(
            fitz.Rect(36, 36, 560, 800),
            "最新简历标准 PDF",
            fontsize=12,
        )
        attach_standard_pdf(latest_resume, document.tobytes())
        document.close()
        client = self.authenticated_client(self.hr)

        response = client.get(
            reverse("talent_pool:detail", args=[membership.pk])
        )

        self.assertContains(
            response,
            reverse("recruitment:preview_resume", args=[latest_resume.pk]),
        )
        self.assertContains(
            response,
            reverse("recruitment:download_resume", args=[latest_resume.pk]),
        )
        self.assertNotContains(
            response,
            reverse(
                "recruitment:preview_resume",
                args=[membership.resume_version.pk],
            ),
        )

    def test_add_tag_recommend_remove_restore_and_purge(self):
        application = self.create_application()
        membership = add_candidate(application.candidate, self.hr)
        same = add_candidate(application.candidate, self.hr)
        self.assertEqual(membership.pk, same.pk)
        self.assertTrue(membership.resume_version.protected)
        tag = TalentTag.objects.create(name="Java", created_by=self.hr)
        TalentTagAssignment.objects.create(
            membership=membership, tag=tag, assigned_by=self.hr
        )
        recommendation, created = recommend_candidate(
            membership, application.position, self.hr
        )
        self.assertTrue(created)
        self.assertEqual(recommendation.source_type, Application.SourceType.TALENT)
        membership.remove(self.hr)
        membership.refresh_from_db()
        self.assertEqual(membership.status, TalentMembership.Status.REMOVED_PENDING)
        membership.restore()
        self.assertEqual(membership.status, TalentMembership.Status.ACTIVE)
        membership.remove(self.hr)
        membership.purge_after = timezone.now() - timedelta(seconds=1)
        membership.save(update_fields=["purge_after"])
        self.assertEqual(purge_removed_memberships(), 1)
        membership.refresh_from_db()
        self.assertEqual(membership.status, TalentMembership.Status.REMOVED)
        self.assertFalse(membership.tag_assignments.exists())

    def test_stale_resume_requires_confirmation(self):
        application = self.create_application()
        membership = add_candidate(application.candidate, self.hr)
        ResumeVersion.objects.filter(pk=membership.resume_version_id).update(
            created_at=timezone.now() - timedelta(days=731)
        )
        membership.resume_version.refresh_from_db()
        with self.assertRaises(TalentPoolError):
            recommend_candidate(
                membership,
                application.position,
                self.hr,
                stale_confirmed=False,
            )

    def test_official_application_links_existing_talent_recommendation(self):
        position = Position.objects.create(
            beisen_position_id="P-LINK",
            name="关联岗位",
        )
        candidate = Candidate.objects.create(applicant_id="A-LINK", name="关联候选人")
        resume, _ = save_resume_bytes(
            candidate, docx_resume_bytes(), "resume.docx"
        )
        membership = add_candidate(candidate, self.hr)
        recommendation, _ = recommend_candidate(
            membership, position, self.hr
        )
        official = upsert_application(
            candidate,
            {
                "applicationId": "APP-LINK",
                "positionId": "P-LINK",
                "positionName": "关联岗位",
            },
        )
        self.assertEqual(official.linked_application, recommendation)

    def test_add_candidate_records_position_and_renders_in_views(self):
        application = self.create_application()
        membership = add_candidate(application.candidate, self.hr)
        self.assertEqual(membership.position, application.position)

        client = self.authenticated_client(self.hr)
        list_res = client.get(reverse("talent_pool:list"))
        self.assertContains(list_res, application.position.name)
        self.assertContains(list_res, "<th>岗位</th>")

        detail_res = client.get(reverse("talent_pool:detail", args=[membership.pk]))
        self.assertContains(detail_res, "来源岗位")
        self.assertContains(detail_res, application.position.name)

    def test_talent_filter_by_position_and_tag(self):
        pos_backend = Position.objects.create(
            beisen_position_id="P-BE",
            name="高级后端开发",
        )
        pos_frontend = Position.objects.create(
            beisen_position_id="P-FE",
            name="前端开发专家",
        )
        app1 = self.create_application(applicant_id="A-BE", position=pos_backend)
        app1.candidate.name = "张后端"
        app1.candidate.save(update_fields=["name"])
        app2 = self.create_application(applicant_id="A-FE", position=pos_frontend)
        app2.candidate.name = "李前端"
        app2.candidate.save(update_fields=["name"])

        mem1 = add_candidate(app1.candidate, self.hr, position=pos_backend)
        mem2 = add_candidate(app2.candidate, self.hr, position=pos_frontend)

        tag = TalentTag.objects.create(name="Go语言", created_by=self.hr)
        TalentTagAssignment.objects.create(membership=mem1, tag=tag, assigned_by=self.hr)

        client = self.authenticated_client(self.hr)

        # Filter by backend position
        res = client.get(reverse("talent_pool:list"), {"position": pos_backend.pk})
        self.assertContains(res, "张后端")
        self.assertNotContains(res, "李前端")

        # Filter by frontend position
        res = client.get(reverse("talent_pool:list"), {"position": pos_frontend.pk})
        self.assertNotContains(res, "张后端")
        self.assertContains(res, "李前端")

        # Filter by backend position + tag
        res = client.get(reverse("talent_pool:list"), {"position": pos_backend.pk, "tag": tag.pk})
        self.assertContains(res, "张后端")

        # Filter by frontend position + tag (no match)
        res = client.get(reverse("talent_pool:list"), {"position": pos_frontend.pk, "tag": tag.pk})
        self.assertNotContains(res, "张后端")
        self.assertNotContains(res, "李前端")
        self.assertContains(res, "暂无人才库成员")

    def test_add_from_application_view_records_position(self):
        application = self.create_application()
        rule_version = PositionRuleVersion.objects.create(
            position=application.position,
            version=1,
            status=PositionRuleVersion.Status.PUBLISHED,
        )
        job = AnalysisJob.objects.create(
            position=application.position,
            requested_by=self.hr,
            status=AnalysisJob.Status.SUCCESS,
        )
        AnalysisItem.objects.create(
            job=job,
            application=application,
            resume_version=application.current_resume,
            rule_version=rule_version,
            status=AnalysisItem.Status.SUCCESS,
        )

        client = self.authenticated_client(self.hr)
        response = client.post(reverse("talent_pool:add_from_application", args=[application.pk]))
        self.assertRedirects(response, reverse("talent_pool:list"))

        membership = TalentMembership.objects.get(candidate=application.candidate)
        self.assertEqual(membership.position, application.position)


class TalentPoolOptimizationTests(WorkflowFixtureMixin, TestCase):
    def test_default_page_size_is_20_and_pagination_navigation(self):
        position = Position.objects.create(beisen_position_id="P-PAG", name="分页测试岗位")
        memberships = []
        for i in range(25):
            candidate = Candidate.objects.create(
                applicant_id=f"A-PAG-{i:03d}",
                name=f"分页候选人-{i:03d}",
                phone=f"1380000{i:04d}",
                email=f"candidate_{i}@example.com",
            )
            save_resume_bytes(candidate, docx_resume_bytes(), f"resume_{i}.docx")
            membership = add_candidate(candidate, self.hr, position=position)
            memberships.append(membership)

        client = self.authenticated_client(self.hr)

        # Page 1
        res1 = client.get(reverse("talent_pool:list"))
        self.assertEqual(res1.status_code, 200)
        self.assertEqual(len(res1.context["memberships"]), 20)
        self.assertEqual(res1.context["paginator"].count, 25)
        self.assertEqual(res1.context["paginator"].num_pages, 2)
        self.assertEqual(res1.context["page_obj"].number, 1)
        self.assertContains(res1, "共 <strong>25</strong> 位成员", count=1)
        self.assertContains(res1, "显示第 1 - 20 位")
        self.assertContains(res1, "下一页")
        self.assertContains(res1, "末页")

        # Page 2
        res2 = client.get(reverse("talent_pool:list"), {"page": 2})
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(len(res2.context["memberships"]), 5)
        self.assertEqual(res2.context["page_obj"].number, 2)
        self.assertContains(res2, "显示第 21 - 25 位")
        self.assertContains(res2, "首页")
        self.assertContains(res2, "上一页")

    def test_invalid_page_fallback(self):
        position = Position.objects.create(beisen_position_id="P-FALLBACK", name="回退测试岗位")
        for i in range(25):
            candidate = Candidate.objects.create(
                applicant_id=f"A-FB-{i:03d}",
                name=f"候选人-{i:03d}",
            )
            save_resume_bytes(candidate, docx_resume_bytes(), f"resume_{i}.docx")
            add_candidate(candidate, self.hr, position=position)

        client = self.authenticated_client(self.hr)

        # Invalid page non-integer
        res_str = client.get(reverse("talent_pool:list"), {"page": "invalid"})
        self.assertEqual(res_str.context["page_obj"].number, 1)

        # Invalid page negative
        res_neg = client.get(reverse("talent_pool:list"), {"page": "-5"})
        self.assertEqual(res_neg.context["page_obj"].number, 1)

        # Out of bounds page (page 9999 on 2-page dataset) -> last page (2)
        res_out = client.get(reverse("talent_pool:list"), {"page": "9999"})
        self.assertEqual(res_out.context["page_obj"].number, 2)

    def test_pagination_preserves_query_filters(self):
        pos1 = Position.objects.create(beisen_position_id="P-PRE1", name="目标开发岗位")
        pos2 = Position.objects.create(beisen_position_id="P-PRE2", name="其他岗位")
        tag = TalentTag.objects.create(name="核心标签", created_by=self.hr)

        for i in range(25):
            candidate = Candidate.objects.create(
                applicant_id=f"A-PRE-{i:03d}",
                name=f"目标候选人-{i:03d}",
            )
            save_resume_bytes(candidate, docx_resume_bytes(), f"resume_{i}.docx")
            mem = add_candidate(candidate, self.hr, position=pos1)
            TalentTagAssignment.objects.create(membership=mem, tag=tag, assigned_by=self.hr)

        client = self.authenticated_client(self.hr)
        res = client.get(
            reverse("talent_pool:list"),
            {"q": "目标", "position": pos1.pk, "tag": tag.pk, "page": 1},
        )
        self.assertEqual(res.status_code, 200)
        self.assertIn("q=%E7%9B%AE%E6%A0%87", res.context["preserved_query"])
        self.assertIn(f"position={pos1.pk}", res.context["preserved_query"])
        self.assertIn(f"tag={tag.pk}", res.context["preserved_query"])
        # Check that page 2 link contains the preserved query
        self.assertContains(res, f"page=2")

    def test_active_and_stale_displayed_removed_excluded(self):
        candidate_active = Candidate.objects.create(applicant_id="A-ACT", name="活跃候选人")
        save_resume_bytes(candidate_active, docx_resume_bytes(), "resume1.docx")
        mem_act = add_candidate(candidate_active, self.hr)

        candidate_stale = Candidate.objects.create(applicant_id="A-STL", name="过期候选人")
        save_resume_bytes(candidate_stale, docx_resume_bytes(), "resume2.docx")
        mem_stl = add_candidate(candidate_stale, self.hr)
        mem_stl.status = TalentMembership.Status.STALE
        mem_stl.save(update_fields=["status"])

        candidate_removed_p = Candidate.objects.create(applicant_id="A-RMP", name="移出待恢复候选人")
        save_resume_bytes(candidate_removed_p, docx_resume_bytes(), "resume3.docx")
        mem_rmp = add_candidate(candidate_removed_p, self.hr)
        mem_rmp.remove(self.hr)

        candidate_removed = Candidate.objects.create(applicant_id="A-RM", name="已移出候选人")
        save_resume_bytes(candidate_removed, docx_resume_bytes(), "resume4.docx")
        mem_rm = add_candidate(candidate_removed, self.hr)
        mem_rm.status = TalentMembership.Status.REMOVED
        mem_rm.save(update_fields=["status"])

        client = self.authenticated_client(self.hr)
        res = client.get(reverse("talent_pool:list"))
        self.assertContains(res, "活跃候选人")
        self.assertContains(res, "过期候选人")
        self.assertNotContains(res, "移出待恢复候选人")
        self.assertNotContains(res, "已移出候选人")

    def test_tag_badges_max_three_and_overflow(self):
        candidate = Candidate.objects.create(applicant_id="A-TAGS", name="多标签候选人")
        save_resume_bytes(candidate, docx_resume_bytes(), "resume.docx")
        membership = add_candidate(candidate, self.hr)

        tags = [
            TalentTag.objects.create(name=f"标签_{i}", created_by=self.hr)
            for i in range(5)
        ]
        for tag in tags:
            TalentTagAssignment.objects.create(membership=membership, tag=tag, assigned_by=self.hr)

        client = self.authenticated_client(self.hr)
        res = client.get(reverse("talent_pool:list"))
        self.assertContains(res, "标签_0")
        self.assertContains(res, "标签_1")
        self.assertContains(res, "标签_2")
        self.assertContains(res, "+2")

        # Detail view shows all 5 tags
        detail_res = client.get(reverse("talent_pool:detail", args=[membership.pk]))
        for tag in tags:
            self.assertContains(detail_res, tag.name)

    def test_tag_list_view_member_count_and_search(self):
        tag_python = TalentTag.objects.create(name="Python开发", created_by=self.hr)
        tag_java = TalentTag.objects.create(name="Java后端", created_by=self.hr)
        tag_empty = TalentTag.objects.create(name="新标签无成员", created_by=self.hr)

        c1 = Candidate.objects.create(applicant_id="A-T1", name="候选人1")
        save_resume_bytes(c1, docx_resume_bytes(), "r1.docx")
        m1 = add_candidate(c1, self.hr)
        TalentTagAssignment.objects.create(membership=m1, tag=tag_python, assigned_by=self.hr)
        TalentTagAssignment.objects.create(membership=m1, tag=tag_java, assigned_by=self.hr)

        c2 = Candidate.objects.create(applicant_id="A-T2", name="候选人2")
        save_resume_bytes(c2, docx_resume_bytes(), "r2.docx")
        m2 = add_candidate(c2, self.hr)
        TalentTagAssignment.objects.create(membership=m2, tag=tag_python, assigned_by=self.hr)

        client = self.authenticated_client(self.hr)
        res = client.get(reverse("talent_pool:tag_list"))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Python开发")
        self.assertContains(res, "<strong>2</strong> 位成员")
        self.assertContains(res, "Java后端")
        self.assertContains(res, "<strong>1</strong> 位成员")
        self.assertContains(res, "新标签无成员")
        self.assertContains(res, "<strong>0</strong> 位成员")

        # Search in tag list
        search_res = client.get(reverse("talent_pool:tag_list"), {"q": "Python"})
        self.assertContains(search_res, "Python开发")
        self.assertNotContains(search_res, "Java后端")
        self.assertNotContains(search_res, "新标签无成员")

    def test_tag_creation_rename_and_permissions(self):
        other_hr = User.objects.create_user(
            username="other_hr",
            password="Other-password-123",
            role=User.Role.HR,
            must_change_password=False,
        )
        tag = TalentTag.objects.create(name="HR专属标签", created_by=self.hr)

        # Non-creator HR cannot rename
        other_client = self.authenticated_client(other_hr)
        rename_res = other_client.post(
            reverse("talent_pool:edit_tag", args=[tag.pk]),
            {"name": "非法重命名"},
            follow=True,
        )
        self.assertContains(rename_res, "只能修改自己创建的标签")
        tag.refresh_from_db()
        self.assertEqual(tag.name, "HR专属标签")

        # Non-creator HR cannot delete
        delete_res = other_client.post(
            reverse("talent_pool:delete_tag", args=[tag.pk]),
            follow=True,
        )
        self.assertContains(delete_res, "只能删除自己创建的标签")
        tag.refresh_from_db()
        self.assertTrue(tag.is_active)

        # Creator HR can rename
        hr_client = self.authenticated_client(self.hr)
        hr_client.post(
            reverse("talent_pool:edit_tag", args=[tag.pk]),
            {"name": "合法重命名标签"},
        )
        tag.refresh_from_db()
        self.assertEqual(tag.name, "合法重命名标签")

        # Admin can delete
        admin_client = self.authenticated_client(self.admin)
        del_res = admin_client.post(
            reverse("talent_pool:delete_tag", args=[tag.pk]),
            follow=True,
        )
        self.assertContains(del_res, "标签已删除")
        tag.refresh_from_db()
        self.assertFalse(tag.is_active)




class DeletionTests(WorkflowFixtureMixin, TestCase):
    def test_hr_can_delete_unanalysed_application(self):
        application = self.create_application()
        client = self.authenticated_client(self.hr)

        detail_response = client.get(
            reverse("recruitment:position_detail", args=[application.position_id])
        )
        candidate_response = client.get(
            reverse("recruitment:candidate_detail", args=[application.candidate_id])
        )
        response = client.post(
            reverse("recruitment:delete_application", args=[application.pk]),
            {"reason": "不再继续处理"},
        )

        self.assertContains(
            detail_response,
            reverse("recruitment:delete_application", args=[application.pk]),
        )
        self.assertContains(
            candidate_response,
            reverse("recruitment:delete_application", args=[application.pk]),
        )
        self.assertRedirects(
            response,
            reverse("recruitment:position_detail", args=[application.position_id]),
        )
        application.refresh_from_db()
        self.assertIsNotNone(application.deleted_at)

    def test_hr_can_confirm_and_bulk_delete_selected_applications(self):
        first = self.create_application(applicant_id="A-BULK-1")
        second = self.create_application(
            applicant_id="A-BULK-2",
            position=first.position,
        )
        client = self.authenticated_client(self.hr)
        url = reverse(
            "recruitment:bulk_delete_applications",
            args=[first.position_id],
        )

        response = client.post(
            url,
            {"application_ids": [first.pk, second.pk]},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "确认批量删除")

        response = client.post(
            url,
            {
                "application_ids": [first.pk, second.pk],
                "confirmed": "1",
                "reason": "批量清理测试投递",
            },
        )

        self.assertRedirects(
            response,
            reverse("recruitment:position_detail", args=[first.position_id]),
        )
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertIsNotNone(first.deleted_at)
        self.assertIsNotNone(second.deleted_at)

    def test_purge_removes_business_data_and_keeps_exclusion_marker(self):
        application = self.create_application()
        self.mark_analyzed(application)
        soft_delete_application(application, self.hr, "人工删除")
        Application.objects.filter(pk=application.pk).update(
            purge_after=timezone.now() - timedelta(seconds=1)
        )
        self.assertEqual(purge_expired_applications(), 1)
        self.assertFalse(Application.objects.filter(pk=application.pk).exists())
        self.assertFalse(
            Candidate.objects.filter(pk=application.candidate_id).exists()
        )
        self.assertTrue(
            ExclusionMarker.objects.filter(application_id=application.application_id).exists()
        )

    def test_purged_application_is_not_recreated_by_sync(self):
        candidate = Candidate.objects.create(applicant_id="A-EXCLUDED")
        ExclusionMarker.objects.create(application_id="APP-EXCLUDED")

        application = upsert_application(
            candidate,
            {
                "applicationId": "APP-EXCLUDED",
                "positionId": "P-EXCLUDED",
                "positionName": "已排除岗位",
            },
        )

        self.assertIsNone(application)
        self.assertFalse(
            Application.objects.filter(application_id="APP-EXCLUDED").exists()
        )

    def test_delete_and_clear_sync_jobs(self):
        client = self.authenticated_client(self.hr)
        finished_job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.MANUAL,
            status=SyncJob.Status.SUCCESS,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=1),
            window_end=timezone.now(),
        )
        running_job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.MANUAL,
            status=SyncJob.Status.RUNNING,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=1),
            window_end=timezone.now(),
        )

        # Cannot delete running job
        res = client.post(reverse("recruitment:delete_sync_job", args=[running_job.pk]))
        self.assertRedirects(res, reverse("recruitment:sync_jobs"))
        self.assertTrue(SyncJob.objects.filter(pk=running_job.pk).exists())

        # Can delete finished job
        res = client.post(reverse("recruitment:delete_sync_job", args=[finished_job.pk]))
        self.assertRedirects(res, reverse("recruitment:sync_jobs"))
        self.assertFalse(SyncJob.objects.filter(pk=finished_job.pk).exists())

        # Create multiple finished jobs and clear all
        for i in range(3):
            SyncJob.objects.create(
                sync_type=SyncJob.SyncType.INCREMENTAL,
                status=SyncJob.Status.SUCCESS,
                window_start=timezone.now() - timedelta(days=1),
                window_end=timezone.now(),
            )
        res = client.post(reverse("recruitment:clear_sync_jobs"))
        self.assertRedirects(res, reverse("recruitment:sync_jobs"))
        self.assertEqual(SyncJob.objects.filter(status=SyncJob.Status.SUCCESS).count(), 0)
        self.assertTrue(SyncJob.objects.filter(pk=running_job.pk).exists())

    def test_delete_and_clear_position_initializations(self):
        client = self.authenticated_client(self.hr)
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.MANUAL,
            status=SyncJob.Status.SUCCESS,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=1),
            window_end=timezone.now(),
        )
        pos1 = Position.objects.create(name="前端开发", beisen_position_id="P-TEST-1")
        pos2 = Position.objects.create(name="后端开发", beisen_position_id="P-TEST-2")
        init1 = PositionRuleInitialization.objects.create(
            sync_job=job,
            position=pos1,
            requested_by=self.hr,
            status=PositionRuleInitialization.Status.SUCCESS,
        )
        init2 = PositionRuleInitialization.objects.create(
            sync_job=job,
            position=pos2,
            requested_by=self.hr,
            status=PositionRuleInitialization.Status.FAILED,
        )

        res = client.post(
            reverse("recruitment:delete_position_initialization", args=[job.pk, init1.pk])
        )
        self.assertRedirects(res, reverse("recruitment:position_initializations", args=[job.pk]))
        self.assertFalse(PositionRuleInitialization.objects.filter(pk=init1.pk).exists())
        self.assertTrue(PositionRuleInitialization.objects.filter(pk=init2.pk).exists())

        res = client.post(
            reverse("recruitment:clear_position_initializations", args=[job.pk])
        )
        self.assertRedirects(res, reverse("recruitment:position_initializations", args=[job.pk]))
        self.assertFalse(PositionRuleInitialization.objects.filter(pk=init2.pk).exists())

    def test_delete_and_clear_notifications(self):
        client = self.authenticated_client(self.hr)
        n1 = self.hr.notifications.create(title="通知1", message="内容1")
        n2 = self.hr.notifications.create(title="通知2", message="内容2")
        admin_n = self.admin.notifications.create(title="管理员通知", message="内容")

        # Delete single notification
        res = client.post(reverse("recruitment:delete_notification", args=[n1.pk]))
        self.assertRedirects(res, reverse("recruitment:notifications"))
        self.assertFalse(self.hr.notifications.filter(pk=n1.pk).exists())

        # Clear all user's notifications
        res = client.post(reverse("recruitment:clear_notifications"))
        self.assertRedirects(res, reverse("recruitment:notifications"))
        self.assertEqual(self.hr.notifications.count(), 0)
        self.assertTrue(self.admin.notifications.filter(pk=admin_n.pk).exists())

    def test_delete_analysis_job_and_review_batch(self):
        client = self.authenticated_client(self.hr)
        pos = Position.objects.create(name="产品经理", beisen_position_id="P-TEST-PM")
        analysis_job = AnalysisJob.objects.create(
            position=pos,
            requested_by=self.hr,
            status=AnalysisJob.Status.SUCCESS,
            total_count=0,
        )
        running_analysis = AnalysisJob.objects.create(
            position=pos,
            requested_by=self.hr,
            status=AnalysisJob.Status.RUNNING,
            total_count=1,
        )

        # Cannot delete running analysis job
        res = client.post(reverse("analysis:job_delete", args=[running_analysis.pk]))
        self.assertTrue(AnalysisJob.objects.filter(pk=running_analysis.pk).exists())

        # Can delete completed analysis job
        res = client.post(reverse("analysis:job_delete", args=[analysis_job.pk]))
        self.assertRedirects(res, reverse("recruitment:position_detail", args=[pos.pk]))
        self.assertFalse(AnalysisJob.objects.filter(pk=analysis_job.pk).exists())

        # Delete review batch
        reviewer = Reviewer.objects.create(name="张主管", email="zhang@example.com")
        batch = ReviewBatch.objects.create(
            position=pos,
            created_by=self.hr,
            reviewer=reviewer,
            status=ReviewBatch.Status.REVOKED,
            expires_at=timezone.now() + timedelta(days=3),
        )
        res = client.post(reverse("reviews:delete", args=[batch.pk]))
        self.assertRedirects(res, reverse("reviews:list"))
        self.assertFalse(ReviewBatch.objects.filter(pk=batch.pk).exists())


class PageSmokeTests(WorkflowFixtureMixin, TestCase):
    def test_active_tasks_expose_cancel_actions(self):
        application = self.create_application(applicant_id="A-CANCEL-PAGE")
        sync_job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.RECONCILIATION,
            status=SyncJob.Status.RUNNING,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
        )
        initialization = PositionRuleInitialization.objects.create(
            sync_job=sync_job,
            position=application.position,
            requested_by=self.hr,
            status=PositionRuleInitialization.Status.QUEUED,
        )
        rule = self.create_rule(application.position)
        analysis_job = AnalysisJob.objects.create(
            position=application.position,
            requested_by=self.hr,
            status=AnalysisJob.Status.RUNNING,
            total_count=1,
        )
        AnalysisItem.objects.create(
            job=analysis_job,
            application=application,
            resume_version=application.current_resume,
            rule_version=rule,
        )
        client = self.authenticated_client(self.hr)

        self.assertContains(
            client.get(reverse("recruitment:sync_jobs")),
            reverse("recruitment:cancel_sync_job", args=[sync_job.pk]),
        )
        self.assertContains(
            client.get(
                reverse(
                    "recruitment:position_initializations",
                    args=[sync_job.pk],
                )
            ),
            reverse(
                "recruitment:cancel_position_initialization",
                args=[sync_job.pk, initialization.pk],
            ),
        )
        self.assertContains(
            client.get(reverse("analysis:job_detail", args=[analysis_job.pk])),
            reverse("analysis:job_cancel", args=[analysis_job.pk]),
        )

        client.post(reverse("recruitment:cancel_sync_job", args=[sync_job.pk]))
        client.post(
            reverse(
                "recruitment:cancel_position_initialization",
                args=[sync_job.pk, initialization.pk],
            )
        )
        client.post(reverse("analysis:job_cancel", args=[analysis_job.pk]))
        sync_job.refresh_from_db()
        initialization.refresh_from_db()
        analysis_job.refresh_from_db()
        self.assertEqual(
            sync_job.status,
            SyncJob.Status.CANCELLATION_REQUESTED,
        )
        self.assertEqual(
            initialization.status,
            PositionRuleInitialization.Status.CANCELLED,
        )
        self.assertEqual(
            analysis_job.status,
            AnalysisJob.Status.CANCELLATION_REQUESTED,
        )

    def test_notification_view_marks_read_and_redirects_safely(self):
        notification = Notification.objects.create(
            user=self.hr,
            title="审核完成",
            target_url="/reviews/123/",
        )
        client = self.authenticated_client(self.hr)

        self.assertContains(
            client.get(reverse("recruitment:notifications")),
            reverse("recruitment:notification_view", args=[notification.pk]),
        )
        self.assertContains(
            client.get(reverse("dashboard")),
            reverse("recruitment:notification_view", args=[notification.pk]),
        )
        response = client.get(
            reverse("recruitment:notification_view", args=[notification.pk])
        )

        self.assertRedirects(response, "/reviews/123/", fetch_redirect_response=False)
        notification.refresh_from_db()
        self.assertIsNotNone(notification.read_at)

        unsafe = Notification.objects.create(
            user=self.hr,
            title="异常目标",
            target_url="//example.com/phishing",
        )
        response = client.get(
            reverse("recruitment:notification_view", args=[unsafe.pk])
        )
        self.assertRedirects(response, reverse("recruitment:notifications"))

        other_users_notification = Notification.objects.create(
            user=self.admin,
            title="管理员通知",
            target_url="/reviews/",
        )
        response = client.get(
            reverse(
                "recruitment:notification_view",
                args=[other_users_notification.pk],
            )
        )
        self.assertEqual(response.status_code, 404)

    @patch("recruitment.views.dispatch_task")
    def test_hr_can_request_missing_beisen_resume(self, dispatch_task):
        application = self.create_application(applicant_id="A-PULL")
        application.current_resume = None
        application.save(update_fields=["current_resume"])
        client = self.authenticated_client(self.hr)

        detail_response = client.get(
            reverse("recruitment:position_detail", args=[application.position_id])
        )
        response = client.post(
            reverse("recruitment:pull_resume", args=[application.pk])
        )

        self.assertContains(detail_response, "按需补拉")
        self.assertRedirects(
            response,
            reverse("recruitment:position_detail", args=[application.position_id]),
        )
        dispatch_task.assert_called_once_with(
            pull_application_resume, application.pk, self.hr.pk
        )

    def test_pull_resume_task_downloads_parses_and_links_resume(self):
        application = self.create_application(applicant_id="A-PULL-TASK")
        application.current_resume.delete()
        application.current_resume = None
        application.save(update_fields=["current_resume"])

        with patch(
            "recruitment.tasks.ITalentClient",
            return_value=FakeITalentClient(),
        ):
            resume_id = pull_application_resume(application.pk, self.hr.pk)

        application.refresh_from_db()
        self.assertEqual(str(application.current_resume_id), resume_id)
        self.assertEqual(
            application.current_resume.parse_status,
            ResumeVersion.ParseStatus.SUCCESS,
        )
        self.assertTrue(application.current_resume.standard_pdf)
        self.assertTrue(
            self.hr.notifications.filter(title="简历补拉完成").exists()
        )

    def test_pull_resume_task_notifies_requester_when_file_is_missing(self):
        application = self.create_application(applicant_id="A-PULL-MISSING")
        application.current_resume.delete()
        application.current_resume = None
        application.save(update_fields=["current_resume"])

        client = FakeITalentClient()
        client.get_resume_file_info = lambda applicant_id, origin=True: {"data": {}}
        with patch("recruitment.tasks.ITalentClient", return_value=client):
            resume_id = pull_application_resume(application.pk, self.hr.pk)

        application.refresh_from_db()
        self.assertEqual(resume_id, "")
        self.assertIsNone(application.current_resume)
        self.assertTrue(
            self.hr.notifications.filter(title="简历补拉失败").exists()
        )

    def test_candidate_without_contact_information_renders_placeholders(self):
        application = self.create_application(applicant_id="A-NO-CONTACT")
        application.candidate.phone = ""
        application.candidate.email = ""
        application.candidate.save(update_fields=["phone", "email"])
        client = self.authenticated_client(self.hr)

        response = client.get(
            reverse("recruitment:candidate_detail", args=[application.candidate_id])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "<dt>手机号</dt><dd>-</dd>", html=True)
        self.assertContains(response, "<dt>邮箱</dt><dd>-</dd>", html=True)

    def test_candidate_detail_displays_synced_resume_modules_and_channel(self):
        application = self.create_application()
        application.source_channel = "招聘官网"
        application.save(update_fields=["source_channel"])
        application.candidate.skills_text = "Python · 熟练"
        application.candidate.resume_modules = {
            "ApplicantEducation": {
                "moduleInfo": [
                    [
                        {"name": "SchoolName", "text": "示例大学"},
                        {"name": "EducationLevel", "text": "本科"},
                        {"name": "MajorName", "text": "计算机科学"},
                        {"name": "StartDate", "text": "2018/09"},
                        {"name": "EndDate", "text": "2022/06"},
                    ]
                ]
            },
            "ApplicantWorkExperience": {
                "moduleInfo": [
                    [
                        {"name": "CompanyName", "text": "示例科技"},
                        {"name": "JobTitle", "text": "后端工程师"},
                        {"name": "StartDate", "text": "2022/07"},
                        {"name": "EndDate", "text": "至今"},
                        {"name": "JobDuty", "text": "负责招聘系统研发。"},
                    ]
                ]
            },
        }
        application.candidate.save(update_fields=["skills_text", "resume_modules"])
        client = self.authenticated_client(self.hr)

        response = client.get(
            reverse("recruitment:candidate_detail", args=[application.candidate_id])
        )

        self.assertContains(response, "教育经历")
        self.assertContains(response, "示例大学")
        self.assertContains(response, "计算机科学")
        self.assertContains(response, "工作经历")
        self.assertContains(response, "示例科技")
        self.assertContains(response, "负责招聘系统研发。")
        self.assertContains(response, "招聘官网")
        self.assertContains(response, "Python · 熟练")

    def test_candidate_detail_uses_business_friendly_labels_and_stacked_experience(self):
        application = self.create_application(applicant_id="56e8dc58-technical-id")
        application.current_resume.original_filename = (
            "0aab56cb1e06458a9b42343861cb9459.pdf"
        )
        application.current_resume.save(update_fields=["original_filename"])
        application.candidate.school = "福建工程学院"
        application.candidate.resume_modules = {
            "ApplicantEducation": {
                "moduleInfo": [
                    [
                        {"name": "SchoolName", "value": "-32767", "text": ""},
                        {
                            "name": "OgSchoolName",
                            "value": "福建工程学院",
                            "text": "福建工程学院",
                        },
                        {"name": "EducationLevel", "value": "1", "text": "本科"},
                    ]
                ]
            }
        }
        application.candidate.save(update_fields=["school", "resume_modules"])
        client = self.authenticated_client(self.hr)

        response = client.get(
            reverse("recruitment:candidate_detail", args=[application.candidate_id])
        )

        self.assertNotContains(response, "Applicant ID")
        self.assertNotContains(response, "56e8dc58-technical-id")
        self.assertNotContains(response, "-32767")
        self.assertNotContains(response, "0aab56cb1e06458a9b42343861cb9459.pdf")
        self.assertNotContains(response, "质量 100.00")
        self.assertContains(response, "测试候选人-后端工程师-简历.pdf")
        self.assertContains(response, "福建工程学院")
        self.assertContains(response, "candidate-experience-stack")

        preview_response = client.get(
            reverse(
                "recruitment:preview_resume",
                args=[application.current_resume_id],
            )
        )
        try:
            self.assertIn(
                "测试候选人-后端工程师-简历.pdf",
                unquote(preview_response.headers["Content-Disposition"]),
            )
        finally:
            preview_response.close()

    def test_position_pages_display_jd_channel_and_visible_application_count(self):
        application = self.create_application()
        application.source_channel = "招聘官网"
        application.save(update_fields=["source_channel"])
        deleted = self.create_application(
            applicant_id="A-DELETED",
            position=application.position,
        )
        deleted.soft_delete(self.hr, "不合适")
        client = self.authenticated_client(self.hr)

        list_response = client.get(reverse("recruitment:position_list"))
        detail_response = client.get(
            reverse("recruitment:position_detail", args=[application.position_id])
        )

        self.assertContains(list_response, "<td>1</td>", html=True)
        self.assertContains(detail_response, "北森岗位 JD")
        self.assertContains(detail_response, "负责后端系统研发。")
        self.assertContains(detail_response, "招聘官网")

    def test_position_detail_uses_sticky_controls_and_internal_table_scroll(self):
        application = self.create_application()
        client = self.authenticated_client(self.hr)

        response = client.get(
            reverse("recruitment:position_detail", args=[application.position_id])
        )

        self.assertContains(response, "position-filter-panel")
        self.assertContains(response, "application-toolbar")
        self.assertContains(response, "application-table-scroll")

    def test_candidate_detail_returns_to_source_position(self):
        application = self.create_application()
        client = self.authenticated_client(self.hr)
        position_url = reverse(
            "recruitment:position_detail", args=[application.position_id]
        )

        position_response = client.get(position_url)
        candidate_url = (
            reverse(
                "recruitment:candidate_detail",
                args=[application.candidate_id],
            )
            + f"?position_id={application.position_id}"
        )
        detail_response = client.get(candidate_url)

        self.assertContains(position_response, candidate_url)
        self.assertContains(detail_response, "返回上一级")
        self.assertContains(detail_response, position_url)
        self.assertContains(
            detail_response,
            f'target="_blank" rel="noopener" href="'
            f'{reverse("recruitment:preview_resume", args=[application.current_resume_id])}"',
        )

    @patch("recruitment.views.dispatch_task")
    def test_hr_can_start_recommended_seven_day_sync(self, dispatch_task):
        client = self.authenticated_client(self.hr)
        before = timezone.now()

        response = client.post(
            reverse("recruitment:sync_jobs"),
            {"sync_type": SyncJob.SyncType.RECONCILIATION},
        )

        self.assertRedirects(response, reverse("recruitment:sync_jobs"))
        job = SyncJob.objects.get()
        self.assertEqual(job.sync_type, SyncJob.SyncType.RECONCILIATION)
        self.assertEqual(job.requested_by, self.hr)
        self.assertLess(abs((job.window_end - before).total_seconds()), 5)
        self.assertLess(
            abs((job.window_end - job.window_start - timedelta(days=7)).total_seconds()),
            1,
        )
        dispatch_task.assert_called_once_with(execute_sync_job, job.pk)

    @patch("recruitment.views.dispatch_task")
    def test_sync_rejects_unknown_range(self, dispatch_task):
        client = self.authenticated_client(self.admin)

        response = client.post(
            reverse("recruitment:sync_jobs"),
            {"sync_type": "unknown"},
        )

        self.assertRedirects(response, reverse("recruitment:sync_jobs"))
        self.assertFalse(SyncJob.objects.exists())
        dispatch_task.assert_not_called()

    def test_sync_page_displays_job_level_error_without_failure_count(self):
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.FULL,
            status=SyncJob.Status.FAILED,
            window_start=timezone.now() - timedelta(days=1),
            window_end=timezone.now(),
            error_message="北森接口返回无法解析的 JSON 内容。",
        )
        client = self.authenticated_client(self.admin)

        response = client.get(reverse("recruitment:sync_jobs"))

        self.assertContains(response, "任务未完成")
        self.assertContains(response, "查看问题")
        self.assertNotContains(response, "无法解析的 JSON")
        detail_response = client.get(
            reverse("recruitment:sync_job_issues", args=[job.pk])
        )
        self.assertContains(detail_response, "同步任务未完成")
        self.assertContains(detail_response, "北森返回的数据暂时无法处理")
        self.assertNotContains(detail_response, "无法解析的 JSON")

    @patch("recruitment.views.dispatch_task")
    def test_sync_issue_page_identifies_candidate_and_can_retry_preview(
        self, dispatch_task
    ):
        application = self.create_application(applicant_id="A-PREVIEW-ISSUE")
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.RECONCILIATION,
            status=SyncJob.Status.PARTIAL,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
            total_count=1,
            success_count=1,
            failure_count=1,
            metadata={
                "candidates": 1,
                "applications": 1,
                "positions": 1,
                "resume_file_updates": 0,
                "record_failures": [],
                "file_failures": [],
                "parse_issues": [],
                "preview_issues": [
                    {
                        "applicant_id": application.candidate.applicant_id,
                        "error": (
                            "[SSL: UNEXPECTED_EOF_WHILE_READING] "
                            "EOF occurred in violation of protocol"
                        ),
                    }
                ],
                "failure_summary": {
                    "record": 0,
                    "file": 0,
                    "parse": 0,
                    "preview": 1,
                },
            },
        )
        client = self.authenticated_client(self.hr)

        list_response = client.get(reverse("recruitment:sync_jobs"))
        issues_url = reverse("recruitment:sync_job_issues", args=[job.pk])
        self.assertContains(list_response, issues_url)
        self.assertContains(list_response, "查看问题")

        response = client.get(issues_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, application.candidate.name)
        self.assertContains(response, application.position.name)
        self.assertContains(response, "简历在线预览文件暂未更新")
        self.assertContains(response, "重新获取预览")
        self.assertNotContains(response, application.candidate.applicant_id)
        self.assertNotContains(response, "UNEXPECTED_EOF_WHILE_READING")

        retry_url = reverse(
            "recruitment:retry_sync_issue",
            args=[job.pk, application.candidate_id],
        )
        retry_response = client.post(retry_url)

        self.assertRedirects(retry_response, issues_url)
        dispatch_task.assert_called_once_with(
            refresh_candidate_resume_preview,
            application.candidate_id,
            self.hr.pk,
        )

    def test_refresh_candidate_resume_preview_updates_existing_resume(self):
        application = self.create_application(applicant_id="A-REFRESH-PREVIEW")
        self.assertFalse(application.current_resume.standard_pdf)

        with patch(
            "recruitment.tasks.ITalentClient",
            return_value=FakeITalentClient(),
        ):
            resume_id = refresh_candidate_resume_preview(
                application.candidate_id,
                self.hr.pk,
            )

        application.refresh_from_db()
        application.current_resume.refresh_from_db()
        self.assertEqual(str(application.current_resume_id), resume_id)
        self.assertTrue(application.current_resume.standard_pdf)
        self.assertTrue(
            self.hr.notifications.filter(title="简历预览更新完成").exists()
        )

    def test_sync_page_auto_refreshes_while_job_is_active(self):
        SyncJob.objects.create(
            sync_type=SyncJob.SyncType.MANUAL,
            status=SyncJob.Status.RUNNING,
            window_start=timezone.now() - timedelta(days=1),
            window_end=timezone.now(),
        )
        client = self.authenticated_client(self.hr)

        response = client.get(reverse("recruitment:sync_jobs"))

        self.assertContains(response, "同步或岗位初始化任务正在执行")
        self.assertNotContains(response, "data-auto-refresh")
        self.assertContains(response, 'data-refresh-region="sync-jobs"')

    def test_authenticated_pages_use_fixed_system_title(self):
        application = self.create_application(applicant_id="A-FIXED-TITLE")
        client = self.authenticated_client(self.hr)

        response = client.get(
            reverse(
                "recruitment:position_detail",
                args=[application.position_id],
            )
        )

        self.assertContains(response, "<title>智筛招聘</title>", html=True)
        self.assertNotContains(response, f"<title>{application.position.name}</title>")

    def test_page_script_preserves_position_and_uses_partial_refresh(self):
        script = (
            Path(settings.BASE_DIR) / "static" / "js" / "app.js"
        ).read_text(encoding="utf-8")

        self.assertIn("savePageState()", script)
        self.assertIn("restorePageState()", script)
        self.assertIn("submitPageForm", script)
        self.assertIn('submitter?.hasAttribute("formaction")', script)
        self.assertNotIn("submitter?.formAction || form.action", script)
        self.assertIn("scheduleRegionRefresh", script)
        self.assertIn("region.replaceWith(nextRegion)", script)
        self.assertNotIn("window.location.reload()", script)
        review_css = (
            Path(settings.BASE_DIR) / "static" / "css" / "app.css"
        ).read_text(encoding="utf-8")
        self.assertIn(".public-review-analysis-scroll", review_css)
        self.assertIn(".public-insight-grid .insight-list", review_css)
        self.assertIn("overflow-x: auto", review_css)
        self.assertIn(".public-review-decision-card .grid-two", review_css)

    @patch("recruitment.views.dispatch_task")
    def test_rule_initialization_progress_is_separate_and_failed_task_can_retry(
        self, dispatch_task
    ):
        application = self.create_application(applicant_id="A-RULE-INIT")
        job = SyncJob.objects.create(
            sync_type=SyncJob.SyncType.FULL,
            status=SyncJob.Status.SUCCESS,
            requested_by=self.hr,
            window_start=timezone.now() - timedelta(days=7),
            window_end=timezone.now(),
            total_count=1,
            success_count=1,
        )
        initialization = PositionRuleInitialization.objects.create(
            sync_job=job,
            position=application.position,
            requested_by=self.hr,
            status=PositionRuleInitialization.Status.FAILED,
            retry_count=1,
            error_message="模型服务暂时不可用。",
            started_at=timezone.now(),
            finished_at=timezone.now(),
        )
        client = self.authenticated_client(self.hr)
        detail_url = reverse(
            "recruitment:position_initializations",
            args=[job.pk],
        )

        list_response = client.get(reverse("recruitment:sync_jobs"))
        detail_response = client.get(detail_url)

        self.assertContains(list_response, "完成 0/1")
        self.assertContains(list_response, detail_url)
        self.assertContains(detail_response, "新岗位规则 V0 独立生成进度")
        self.assertContains(detail_response, "模型服务暂时不可用")
        self.assertContains(detail_response, "重新生成")

        response = client.post(
            reverse(
                "recruitment:retry_position_initialization",
                args=[job.pk, initialization.pk],
            )
        )

        self.assertRedirects(response, detail_url)
        initialization.refresh_from_db()
        self.assertEqual(
            initialization.status,
            PositionRuleInitialization.Status.QUEUED,
        )
        self.assertEqual(initialization.error_message, "")
        self.assertIsNone(initialization.started_at)
        self.assertIsNone(initialization.finished_at)
        dispatch_task.assert_called_once_with(
            execute_position_rule_initialization,
            initialization.pk,
        )

    def test_primary_authenticated_pages_render(self):
        application = self.create_application()
        self.mark_analyzed(application)
        membership = add_candidate(application.candidate, self.hr)
        hr_client = self.authenticated_client(self.hr)
        admin_client = self.authenticated_client(self.admin)
        hr_urls = [
            reverse("dashboard"),
            reverse("recruitment:position_list"),
            reverse("recruitment:position_detail", args=[application.position_id]),
            reverse("recruitment:candidate_detail", args=[application.candidate_id]),
            reverse("reviews:list"),
            reverse("talent_pool:list"),
            reverse("talent_pool:detail", args=[membership.pk]),
            reverse("recruitment:sync_jobs"),
        ]
        admin_urls = [
            reverse("analysis:rule_list"),
            reverse("analysis:usage"),
            reverse("recruitment:recycle_bin"),
            reverse("accounts:user_list"),
        ]
        for url in hr_urls:
            self.assertEqual(hr_client.get(url).status_code, 200, url)
        for url in admin_urls:
            self.assertEqual(admin_client.get(url).status_code, 200, url)

    def test_resume_download_requires_hr_login(self):
        application = self.create_application(applicant_id="A-DOWNLOAD")
        url = reverse(
            "recruitment:download_resume",
            args=[application.current_resume_id],
        )
        self.assertEqual(self.client.get(url).status_code, 302)
        response = self.authenticated_client(self.hr).get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        response.close()

    def test_preview_missing_physical_file_renders_friendly_page(self):
        application = self.create_application(applicant_id="A-MISSING-FILE")
        resume = application.current_resume
        # Point to a non-existent file name in storage
        resume.standard_pdf.name = "resumes/non_existent_file.pdf"
        resume.save(update_fields=["standard_pdf"])

        client = self.authenticated_client(self.hr)
        url = reverse("recruitment:preview_resume", args=[resume.pk])
        response = client.get(url)
        # Should render 404 with friendly explanation and retry button, never unhandled 500
        self.assertEqual(response.status_code, 404)
        self.assertContains(response, "简历在线预览暂不可用", status_code=404)
        self.assertContains(response, "重新拉取简历", status_code=404)

    def test_public_review_has_top_return_button(self):
        model, _ = ModelVersion.objects.get_or_create(
            name="gpt-5.6-sol", defaults={"provider": "custom", "is_active": True}
        )
        prompt, _ = PromptVersion.objects.get_or_create(
            version=1, defaults={"content": "{}", "is_active": True}
        )
        application = self.create_application(applicant_id="A-REV-TOP")
        analysis_item = self.mark_analyzed(application)
        AnalysisReport.objects.create(
            item=analysis_item,
            score=90,
            rating=AnalysisReport.Rating.PRIORITY,
            model_version=model,
            prompt_version=prompt,
        )
        reviewer = Reviewer.objects.create(name="负责人A", email="rev_a@example.com")
        batch = create_review_batch(
            application.position,
            [application.pk],
            reviewer,
            self.hr,
            72,
        )
        item = batch.items.get()
        token = token_for_batch(batch)
        url = reverse("reviews:public_item", args=[batch.public_id, token, item.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "返回审核列表")
        self.assertContains(response, "data-review-list-return")

    def test_report_detail_has_print_and_export_buttons(self):
        model, _ = ModelVersion.objects.get_or_create(
            name="gpt-5.6-sol", defaults={"provider": "custom", "is_active": True}
        )
        prompt, _ = PromptVersion.objects.get_or_create(
            version=1, defaults={"content": "{}", "is_active": True}
        )
        application = self.create_application(applicant_id="A-REPORT-PRINT")
        item = self.mark_analyzed(application)
        report = AnalysisReport.objects.create(
            item=item,
            score=90,
            rating=AnalysisReport.Rating.PRIORITY,
            model_version=model,
            prompt_version=prompt,
        )
        client = self.authenticated_client(self.hr)
        url = reverse("analysis:report_detail", args=[report.pk])
        response = client.get(url)
        self.assertContains(response, "下载 PDF")
        self.assertNotContains(response, "打印 / 保存为 PDF")

    def test_talent_detail_stacked_layout_and_fields(self):
        candidate = Candidate.objects.create(
            applicant_id="A-TALENT-DETAIL-TEST",
            name="李阳",
            phone="19859813919",
            email="liyang@example.com",
            current_company="福州鼎盛星航贸易有限公司",
            school="福建商学院",
            skills_text="Excel精通, Fastmoss熟练, GMV调控",
            profile={"Age": 26, "NativePlace": "福建福州"},
            resume_modules={
                "ApplicantEducation": {
                    "moduleInfo": [
                        [
                            {"name": "OgSchoolName", "value": "福建商学院", "text": "福建商学院"},
                            {"name": "EducationLevel", "value": "本科", "text": "本科"},
                            {"name": "MajorName", "value": "电子商务", "text": "电子商务"},
                        ]
                    ]
                },
                "ApplicantWorkExperience": {
                    "moduleInfo": [
                        [
                            {"name": "CompanyName", "value": "福州鼎盛星航贸易有限公司", "text": "福州鼎盛星航贸易有限公司"},
                            {"name": "JobTitle", "value": "运营主管", "text": "运营主管"},
                            {"name": "StartDate", "value": "2023-06", "text": "2023-06"},
                            {"name": "EndDate", "value": "至今", "text": "至今"},
                            {"name": "JobDuty", "value": "负责TikTok店铺整体GMV运营", "text": "负责TikTok店铺整体GMV运营"},
                        ]
                    ]
                },
            },
        )
        membership = add_candidate(candidate, self.hr)
        client = self.authenticated_client(self.hr)
        url = reverse("talent_pool:detail", args=[membership.pk])
        response = client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "李阳")
        self.assertContains(response, "26 岁")
        self.assertContains(response, "福建福州")
        self.assertContains(response, "福建商学院")
        self.assertContains(response, "工作经历")
        self.assertContains(response, "福州鼎盛星航贸易有限公司")
        self.assertContains(response, "运营主管")
        self.assertContains(response, "专业技能")
        self.assertContains(response, "Excel精通")
        self.assertContains(response, "团队标签")
        self.assertContains(response, "人才库备注")
        # Ensure separate education section is omitted as requested
        self.assertNotContains(response, "<h2>🎓 教育经历</h2>")
        self.assertNotContains(response, "<h2>教育经历</h2>")

    def test_talent_interview_workflow_and_api(self):
        from datetime import date

        position = Position.objects.create(name="运营助理", position_type="全职")
        candidate = Candidate.objects.create(
            applicant_id="A-INTV-TEST-1",
            name="魏辛真",
            phone="13900000000",
            email="wei@example.com",
        )
        Application.objects.create(
            candidate=candidate,
            position=position,
            source_type=Application.SourceType.TALENT,
            source_channel="BOSS直聘",
        )
        # Adding to talent pool auto-creates TalentInterview
        membership = add_candidate(candidate, self.hr, position=position)
        interview = TalentInterview.objects.get(candidate=candidate)
        self.assertEqual(interview.position_name, "运营助理")
        self.assertEqual(interview.channel, "BOSS直聘")
        self.assertEqual(interview.result, "未面试")

        # Test date and weekday
        interview.interview_date = date(2026, 8, 20)  # Thursday
        interview.interview_time = "09:30"
        interview.first_interviewer = "发添"
        interview.result = "录用"
        interview.notes = "初试通过，表现优秀"
        interview.save()
        self.assertIn("2026年8月20日 星期四", interview.formatted_date_with_weekday)
        self.assertEqual(interview.result_color_type, "success")

        client = self.authenticated_client(self.hr)

        # 1. Test talent_list buttons
        list_resp = client.get(reverse("talent_pool:list"))
        self.assertEqual(list_resp.status_code, 200)
        self.assertContains(list_resp, "面试信息")
        self.assertContains(list_resp, "标签管理")

        # 2. Test interview_list view
        intv_resp = client.get(reverse("talent_pool:interview_list"))
        self.assertEqual(intv_resp.status_code, 200)
        self.assertContains(intv_resp, "魏辛真")
        self.assertContains(intv_resp, "2026年8月20日 星期四")
        self.assertContains(intv_resp, "09:30")
        self.assertContains(intv_resp, "运营助理")
        self.assertContains(intv_resp, "发添")
        self.assertContains(intv_resp, "录用")
        self.assertContains(intv_resp, "BOSS直聘")

        # 3. Test interview_update_api
        update_url = reverse("talent_pool:interview_update", args=[interview.pk])
        update_resp = client.post(
            update_url,
            data=json.dumps(
                {
                    "first_interviewer": "张倩",
                    "second_interviewer": "李总",
                    "result": "待入职",  # Custom result option
                    "notes": "已发放Offer",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(update_resp.status_code, 200)
        json_data = update_resp.json()
        self.assertTrue(json_data["ok"])
        self.assertEqual(json_data["interview"]["first_interviewer"], "张倩")
        self.assertEqual(json_data["interview"]["second_interviewer"], "李总")
        self.assertEqual(json_data["interview"]["result"], "待入职")

        # Check that custom result option is dynamically persisted
        self.assertTrue(InterviewResultOption.objects.filter(name="待入职").exists())

        # 4. Test delete
        del_url = reverse("talent_pool:interview_delete", args=[interview.pk])
        del_resp = client.post(del_url)
        self.assertEqual(del_resp.status_code, 302)
        self.assertFalse(TalentInterview.objects.filter(pk=interview.pk).exists())


